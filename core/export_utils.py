"""
Utilidades para exportación de reportes a Excel (.xlsx) y CSV.
Desarrollado para el Sistema Contable ContaSys - FIIS UNI.
"""
import csv
import io
from decimal import Decimal
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from .models import AsientoContable, CuentaContable, Movimiento
from .reporte_utils import get_reporte_context


# ─── ESTILOS PARA EXCEL ───────────────────────────────────────────────────────

HEADER_FILL = PatternFill(start_color="1F3C88", end_color="1F3C88", fill_type="solid")
SUBHEADER_FILL = PatternFill(start_color="2980B9", end_color="2980B9", fill_type="solid")
TOTAL_FILL = PatternFill(start_color="EBF3FB", end_color="EBF3FB", fill_type="solid")
ACCENT_FILL = PatternFill(start_color="F2F4F8", end_color="F2F4F8", fill_type="solid")

FONT_TITLE = Font(name="Calibri", size=14, bold=True, color="1F3C88")
FONT_SUBTITLE = Font(name="Calibri", size=10, italic=True, color="555555")
FONT_HEADER = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
FONT_BOLD = Font(name="Calibri", size=11, bold=True)
FONT_NORMAL = Font(name="Calibri", size=11)

ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")

BORDER_THIN = Side(border_style="thin", color="CCCCCC")
BORDER_DOUBLE = Side(border_style="double", color="333333")
CELL_BORDER = Border(left=BORDER_THIN, right=BORDER_THIN, top=BORDER_THIN, bottom=BORDER_THIN)
TOTAL_BORDER = Border(top=BORDER_THIN, bottom=BORDER_DOUBLE)

FORMAT_CURRENCY = '"S/" #,##0.00'


def _ajustar_anchos(ws):
    """Autoajusta el ancho de las columnas con un margen de holgura."""
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or '')
            if len(val) > max_len:
                max_len = len(val)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)


# ─── 1. EXCEL: LIBRO DIARIO ───────────────────────────────────────────────────

def build_excel_diario(wb, title_sheet="Libro Diario"):
    ws = wb.active if "Sheet" in wb.sheetnames else wb.create_sheet(title=title_sheet)
    ws.title = title_sheet

    ws.append(["SISTEMA CONTABLE CONTASYS - LIBRO DIARIO"])
    ws.append(["Registro cronológico de operaciones contables"])
    ws.append([])

    ws.cell(1, 1).font = FONT_TITLE
    ws.cell(2, 1).font = FONT_SUBTITLE

    headers = ["N° Asiento", "Fecha", "Glosa / Descripción", "Código", "Cuenta Contable", "Debe (S/)", "Haber (S/)"]
    ws.append(headers)
    header_row = 4

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=header_row, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = FONT_HEADER
        cell.alignment = ALIGN_CENTER if col_idx in (1, 2, 4) else (ALIGN_RIGHT if col_idx in (6, 7) else ALIGN_LEFT)

    asientos = AsientoContable.objects.prefetch_related('movimientos__cuenta').order_by('fecha', 'id')
    current_row = 5
    total_debe = Decimal('0')
    total_haber = Decimal('0')

    for num, asiento in enumerate(asientos, 1):
        for mov in asiento.movimientos.all():
            monto_debe = float(mov.monto) if mov.tipo == 'debe' else 0.0
            monto_haber = float(mov.monto) if mov.tipo == 'haber' else 0.0
            
            if mov.tipo == 'debe':
                total_debe += mov.monto
            else:
                total_haber += mov.monto

            ws.append([
                num,
                asiento.fecha.strftime("%d/%m/%Y"),
                asiento.descripcion,
                mov.cuenta.codigo,
                mov.cuenta.nombre,
                monto_debe,
                monto_haber
            ])

            for c in range(1, 8):
                cell = ws.cell(row=current_row, column=c)
                cell.font = FONT_NORMAL
                cell.border = CELL_BORDER
                if c in (1, 2, 4):
                    cell.alignment = ALIGN_CENTER
                elif c in (6, 7):
                    cell.alignment = ALIGN_RIGHT
                    cell.number_format = FORMAT_CURRENCY
                else:
                    cell.alignment = ALIGN_LEFT

            current_row += 1

    # Fila de totales
    ws.append(["", "", "", "", "TOTAL GENERAL", float(total_debe), float(total_haber)])
    for c in range(1, 8):
        cell = ws.cell(row=current_row, column=c)
        cell.font = FONT_BOLD
        cell.fill = TOTAL_FILL
        cell.border = TOTAL_BORDER
        if c in (6, 7):
            cell.alignment = ALIGN_RIGHT
            cell.number_format = FORMAT_CURRENCY
        elif c == 5:
            cell.alignment = ALIGN_RIGHT

    _ajustar_anchos(ws)
    return ws


# ─── 2. EXCEL: BALANCE DE COMPROBACIÓN ────────────────────────────────────────

def build_excel_balance_comprobacion(wb, title_sheet="Balance Comprobación"):
    ws = wb.create_sheet(title=title_sheet) if title_sheet in wb.sheetnames or wb.active.title != "Sheet" else wb.active
    if ws.title != title_sheet:
        ws = wb.create_sheet(title=title_sheet)

    ws.append(["SISTEMA CONTABLE CONTASYS - BALANCE DE COMPROBACIÓN"])
    ws.append(["Sumas del Mayor y Saldos de Cuentas"])
    ws.append([])

    ws.cell(1, 1).font = FONT_TITLE
    ws.cell(2, 1).font = FONT_SUBTITLE

    # Encabezados en 2 niveles
    ws.append(["Código", "Cuenta Contable", "SUMAS DEL MAYOR", "", "SALDOS", ""])
    ws.append(["", "", "Debe (S/)", "Haber (S/)", "Deudor (S/)", "Acreedor (S/)"])
    
    ws.merge_cells("A4:A5")
    ws.merge_cells("B4:B5")
    ws.merge_cells("C4:D4")
    ws.merge_cells("E4:F4")

    for r in range(4, 6):
        for c in range(1, 7):
            cell = ws.cell(row=r, column=c)
            cell.fill = HEADER_FILL
            cell.font = FONT_HEADER
            cell.alignment = ALIGN_CENTER
            cell.border = CELL_BORDER

    ctx = get_reporte_context()
    datos = ctx.get('bal_comp_datos', [])
    current_row = 6

    for item in datos:
        c = item['cuenta']
        ws.append([
            c.codigo,
            c.nombre,
            float(item['total_debe']),
            float(item['total_haber']),
            float(item['saldo_deudor']),
            float(item['saldo_acreedor']),
        ])
        for col_idx in range(1, 7):
            cell = ws.cell(row=current_row, column=col_idx)
            cell.font = FONT_NORMAL
            cell.border = CELL_BORDER
            if col_idx == 1:
                cell.alignment = ALIGN_CENTER
            elif col_idx == 2:
                cell.alignment = ALIGN_LEFT
            else:
                cell.alignment = ALIGN_RIGHT
                cell.number_format = FORMAT_CURRENCY
        current_row += 1

    # Fila de Totales
    ws.append([
        "TOTAL",
        "SUMAS Y SALDOS CUADRADOS",
        float(ctx.get('gran_total_debe', 0)),
        float(ctx.get('gran_total_haber', 0)),
        float(ctx.get('gran_saldo_deudor', 0)),
        float(ctx.get('gran_saldo_acreedor', 0)),
    ])
    for col_idx in range(1, 7):
        cell = ws.cell(row=current_row, column=col_idx)
        cell.font = FONT_BOLD
        cell.fill = TOTAL_FILL
        cell.border = TOTAL_BORDER
        if col_idx in (1, 2):
            cell.alignment = ALIGN_LEFT if col_idx == 2 else ALIGN_CENTER
        else:
            cell.alignment = ALIGN_RIGHT
            cell.number_format = FORMAT_CURRENCY

    _ajustar_anchos(ws)
    return ws


# ─── 3. EXCEL: ESTADO DE RESULTADOS ───────────────────────────────────────────

def build_excel_estado_resultados(wb, title_sheet="Estado de Resultados"):
    ws = wb.create_sheet(title=title_sheet)
    ws.append(["SISTEMA CONTABLE CONTASYS - ESTADO DE RESULTADOS"])
    ws.append(["Estado Financiero de Rendimiento Económico"])
    ws.append([])

    ws.cell(1, 1).font = FONT_TITLE
    ws.cell(2, 1).font = FONT_SUBTITLE

    ws.append(["CONCEPTO", "IMPORTE (S/)"])
    ws.cell(4, 1).fill = HEADER_FILL
    ws.cell(4, 1).font = FONT_HEADER
    ws.cell(4, 2).fill = HEADER_FILL
    ws.cell(4, 2).font = FONT_HEADER
    ws.cell(4, 2).alignment = ALIGN_RIGHT

    ctx = get_reporte_context()
    lineas = [
        ("Ventas Netas (Ingresos Operacionales)", float(ctx.get('er_total_ventas', 0)), False),
        ("(-) Costo de Ventas", -float(ctx.get('er_total_costo_ventas', 0)), False),
        ("(=) UTILIDAD BRUTA", float(ctx.get('er_utilidad_bruta', 0)), True),
        ("(-) Gastos Operativos (Adm. y Ventas)", -float(ctx.get('er_total_gastos_operativos', 0)), False),
        ("(=) UTILIDAD OPERATIVA", float(ctx.get('er_utilidad_operativa', 0)), True),
        ("(-) Gastos Financieros", -float(ctx.get('er_total_gastos_financieros', 0)), False),
        ("(+) Otros Ingresos", float(ctx.get('er_total_otros_ingresos', 0)), False),
        ("(-) Otros Gastos", -float(ctx.get('er_total_otros_gastos', 0)), False),
        ("(=) RESULTADO ANTES DE IMPUESTOS", float(ctx.get('er_utilidad_antes_impuesto', 0)), True),
        ("(-) Impuesto a la Renta Estimado (30%)", -float(ctx.get('er_impuesto', 0)), False),
        ("(=) UTILIDAD NETA DEL EJERCICIO", float(ctx.get('er_utilidad_neta', 0)), True),
    ]

    r = 5
    for concepto, monto, es_total in lineas:
        ws.append([concepto, monto])
        c1 = ws.cell(row=r, column=1)
        c2 = ws.cell(row=r, column=2)
        c2.number_format = FORMAT_CURRENCY
        c2.alignment = ALIGN_RIGHT

        if es_total:
            c1.font = FONT_BOLD
            c2.font = FONT_BOLD
            c1.fill = TOTAL_FILL
            c2.fill = TOTAL_FILL
            c1.border = TOTAL_BORDER
            c2.border = TOTAL_BORDER
        else:
            c1.font = FONT_NORMAL
            c2.font = FONT_NORMAL
            c1.border = CELL_BORDER
            c2.border = CELL_BORDER
        r += 1

    _ajustar_anchos(ws)
    return ws


# ─── 4. EXCEL: BALANCE GENERAL ────────────────────────────────────────────────

def build_excel_balance_general(wb, title_sheet="Balance General"):
    ws = wb.create_sheet(title=title_sheet)
    ws.append(["SISTEMA CONTABLE CONTASYS - BALANCE GENERAL"])
    ws.append(["Estado de Situación Financiera"])
    ws.append([])

    ws.cell(1, 1).font = FONT_TITLE
    ws.cell(2, 1).font = FONT_SUBTITLE

    ws.append(["CATEGORÍA", "CÓDIGO", "CUENTA CONTABLE", "SALDO (S/)"])
    for c in range(1, 5):
        cell = ws.cell(row=4, column=c)
        cell.fill = HEADER_FILL
        cell.font = FONT_HEADER
        cell.alignment = ALIGN_CENTER if c in (1, 2) else (ALIGN_RIGHT if c == 4 else ALIGN_LEFT)

    ctx = get_reporte_context()
    current_row = 5

    def escribir_seccion(titulo, items, total_seccion):
        nonlocal current_row
        for it in items:
            ws.append([titulo, it['cuenta'].codigo, it['cuenta'].nombre, float(it['saldo'])])
            for c in range(1, 5):
                cell = ws.cell(row=current_row, column=c)
                cell.font = FONT_NORMAL
                cell.border = CELL_BORDER
                if c == 4:
                    cell.alignment = ALIGN_RIGHT
                    cell.number_format = FORMAT_CURRENCY
            current_row += 1
        
        # Subtotal
        ws.append(["", "", f"TOTAL {titulo.upper()}", float(total_seccion)])
        for c in range(1, 5):
            cell = ws.cell(row=current_row, column=c)
            cell.font = FONT_BOLD
            cell.fill = TOTAL_FILL
            cell.border = TOTAL_BORDER
            if c == 4:
                cell.alignment = ALIGN_RIGHT
                cell.number_format = FORMAT_CURRENCY
        current_row += 1

    escribir_seccion("Activo", ctx.get('bg_activos', []), ctx.get('bg_total_activos', 0))
    escribir_seccion("Pasivo", ctx.get('bg_pasivos', []), ctx.get('bg_total_pasivos', 0))
    escribir_seccion("Patrimonio", ctx.get('bg_patrimonio', []), ctx.get('bg_total_patrimonio', 0))

    # Total Pasivo + Patrimonio
    ws.append(["", "", "TOTAL PASIVO + PATRIMONIO", float(ctx.get('bg_total_pasivos', 0) + ctx.get('bg_total_patrimonio', 0))])
    for c in range(1, 5):
        cell = ws.cell(row=current_row, column=c)
        cell.font = Font(name="Calibri", size=11, bold=True, color="1F3C88")
        cell.fill = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
        cell.border = TOTAL_BORDER
        if c == 4:
            cell.alignment = ALIGN_RIGHT
            cell.number_format = FORMAT_CURRENCY

    _ajustar_anchos(ws)
    return ws


# ─── GENERADOR DE EXCEL COMPLETO (MULTI-HOJA) ─────────────────────────────────

def generar_excel_completo():
    wb = Workbook()
    # Eliminar hoja por defecto vacía o usarla
    build_excel_diario(wb, title_sheet="Libro Diario")
    build_excel_balance_comprobacion(wb, title_sheet="Balance de Comprobación")
    build_excel_estado_resultados(wb, title_sheet="Estado de Resultados")
    build_excel_balance_general(wb, title_sheet="Balance General")

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()


# ─── GENERADORES CSV INDIVIDUALES ─────────────────────────────────────────────

def generar_csv_diario():
    output = io.StringIO()
    # UTF-8 BOM para apertura perfecta en Excel en español
    writer = csv.writer(output, delimiter=';')
    writer.writerow(["Nro Asiento", "Fecha", "Descripcion / Glosa", "Codigo Cuenta", "Nombre Cuenta", "Debe", "Haber"])

    asientos = AsientoContable.objects.prefetch_related('movimientos__cuenta').order_by('fecha', 'id')
    for num, a in enumerate(asientos, 1):
        for m in a.movimientos.all():
            debe = f"{m.monto:.2f}" if m.tipo == 'debe' else "0.00"
            haber = f"{m.monto:.2f}" if m.tipo == 'haber' else "0.00"
            writer.writerow([num, a.fecha.strftime("%d/%m/%Y"), a.descripcion, m.cuenta.codigo, m.cuenta.nombre, debe, haber])

    return output.getvalue().encode('utf-8-sig')


def generar_csv_balance_comprobacion():
    output = io.StringIO()
    writer = csv.writer(output, delimiter=';')
    writer.writerow(["Codigo", "Cuenta Contable", "Total Debe", "Total Haber", "Saldo Deudor", "Saldo Acreedor"])

    ctx = get_reporte_context()
    for item in ctx.get('bal_comp_datos', []):
        c = item['cuenta']
        writer.writerow([
            c.codigo,
            c.nombre,
            f"{item['total_debe']:.2f}",
            f"{item['total_haber']:.2f}",
            f"{item['saldo_deudor']:.2f}",
            f"{item['saldo_acreedor']:.2f}"
        ])

    writer.writerow([
        "TOTALES",
        "CUADRE GENERAL",
        f"{ctx.get('gran_total_debe', 0):.2f}",
        f"{ctx.get('gran_total_haber', 0):.2f}",
        f"{ctx.get('gran_saldo_deudor', 0):.2f}",
        f"{ctx.get('gran_saldo_acreedor', 0):.2f}"
    ])
    return output.getvalue().encode('utf-8-sig')


def generar_csv_estado_resultados():
    output = io.StringIO()
    writer = csv.writer(output, delimiter=';')
    writer.writerow(["Concepto", "Monto (PEN)"])

    ctx = get_reporte_context()
    lineas = [
        ("Ventas Netas", f"{ctx.get('er_total_ventas', 0):.2f}"),
        ("(-) Costo de Ventas", f"-{ctx.get('er_total_costo_ventas', 0):.2f}"),
        ("(=) UTILIDAD BRUTA", f"{ctx.get('er_utilidad_bruta', 0):.2f}"),
        ("(-) Gastos Operativos", f"-{ctx.get('er_total_gastos_operativos', 0):.2f}"),
        ("(=) UTILIDAD OPERATIVA", f"{ctx.get('er_utilidad_operativa', 0):.2f}"),
        ("(-) Gastos Financieros", f"-{ctx.get('er_total_gastos_financieros', 0):.2f}"),
        ("(+) Otros Ingresos", f"{ctx.get('er_total_otros_ingresos', 0):.2f}"),
        ("(-) Otros Gastos", f"{-ctx.get('er_total_otros_gastos', 0):.2f}"),
        ("(=) RESULTADO ANTES DE IMPUESTO", f"{ctx.get('er_utilidad_antes_impuesto', 0):.2f}"),
        ("(-) Impuesto a la Renta (30%)", f"-{ctx.get('er_impuesto', 0):.2f}"),
        ("(=) UTILIDAD NETA DEL EJERCICIO", f"{ctx.get('er_utilidad_neta', 0):.2f}"),
    ]
    for c, m in lineas:
        writer.writerow([c, m])

    return output.getvalue().encode('utf-8-sig')


def generar_csv_balance_general():
    output = io.StringIO()
    writer = csv.writer(output, delimiter=';')
    writer.writerow(["Categoria", "Codigo", "Cuenta", "Saldo (PEN)"])

    ctx = get_reporte_context()
    for item in ctx.get('bg_activos', []):
        writer.writerow(["Activo", item['cuenta'].codigo, item['cuenta'].nombre, f"{item['saldo']:.2f}"])
    writer.writerow(["Activo", "TOTAL", "TOTAL ACTIVOS", f"{ctx.get('bg_total_activos', 0):.2f}"])

    for item in ctx.get('bg_pasivos', []):
        writer.writerow(["Pasivo", item['cuenta'].codigo, item['cuenta'].nombre, f"{item['saldo']:.2f}"])
    writer.writerow(["Pasivo", "TOTAL", "TOTAL PASIVOS", f"{ctx.get('bg_total_pasivos', 0):.2f}"])

    for item in ctx.get('bg_patrimonio', []):
        writer.writerow(["Patrimonio", item['cuenta'].codigo, item['cuenta'].nombre, f"{item['saldo']:.2f}"])
    writer.writerow(["Patrimonio", "TOTAL", "TOTAL PATRIMONIO", f"{ctx.get('bg_total_patrimonio', 0):.2f}"])

    total_pp = ctx.get('bg_total_pasivos', 0) + ctx.get('bg_total_patrimonio', 0)
    writer.writerow(["Ecuacion", "TOTAL", "TOTAL PASIVO + PATRIMONIO", f"{total_pp:.2f}"])

    return output.getvalue().encode('utf-8-sig')


# ─── 4. EXCEL: CASO COMPLETO EN HOJA ÚNICA (DIARIO, CUENTAS T Y ESTADOS FINANCIEROS) ───

def generar_excel_caso_hoja_unica(operaciones_datos=None, titulo_caso="CASO CICLO CONTABLE"):
    """
    Genera un archivo Excel con todo el caso contable resuelto en UNA SOLA HOJA para CUALQUIER caso X:
    1. Registro de los diarios en la base de datos (Libro Diario dinámico)
    2. Registro en los mayores (Cuentas T con saldos netos destacados en amarillo)
    3. Estados Financieros (Estado de Situación Financiera y Estado de Resultados calculados dinámicamente)
    
    Si operaciones_datos es proporcionado (lista de dicts con fecha, glosa, movimientos), se genera
    directamente desde la propuesta de la IA en memoria. Si es None, se lee de la base de datos.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Caso Ciclo Contable"

    # Si no se pasan operaciones en memoria, leerlas de la base de datos
    if operaciones_datos is None:
        asientos_qs = AsientoContable.objects.prefetch_related('movimientos__cuenta').order_by('fecha', 'id')
        operaciones_datos = []
        for a in asientos_qs:
            movs = []
            for m in a.movimientos.all():
                movs.append({
                    'codigo_cuenta': m.cuenta.codigo,
                    'nombre_cuenta': m.cuenta.nombre,
                    'tipo': m.tipo,
                    'monto': float(m.monto)
                })
            operaciones_datos.append({
                'fecha': a.fecha.strftime("%d/%m/%Y"),
                'glosa': a.descripcion,
                'movimientos': movs
            })

    # Estilos específicos para la hoja única
    YELLOW_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    GREEN_ACCENT = PatternFill(start_color="D1E7DD", end_color="D1E7DD", fill_type="solid")
    BLUE_HEADER = PatternFill(start_color="1F3C88", end_color="1F3C88", fill_type="solid")
    TOTAL_FILL = PatternFill(start_color="EBF3FB", end_color="EBF3FB", fill_type="solid")
    THIN_BORDER = Border(left=BORDER_THIN, right=BORDER_THIN, top=BORDER_THIN, bottom=BORDER_THIN)

    # Encabezado principal dinámico
    ws.cell(2, 2, (titulo_caso or "MINI CASO CICLO CONTABLE").upper()).font = Font(name="Calibri", size=16, bold=True, color="1F3C88")
    ws.cell(3, 2, "Universidad Nacional de Ingeniería - FIIS | Sistema y Gestión Financiera (GE605)").font = FONT_SUBTITLE

    # ═════════════════════════════════════════════════════════════════════════
    # 1. REGISTRO DE LOS DIARIOS EN LA BASE DE DATOS
    # ═════════════════════════════════════════════════════════════════════════
    r = 5
    ws.cell(r, 2, "1. Registro de los diarios en la base de datos:").font = Font(name="Calibri", size=12, bold=True, color="1F3C88")
    r += 1

    headers_diario = ["Asiento", "Fecha", "Código", "Cuenta Contable", "DEBE (S/)", "HABER (S/)", "Glosa / Descripción"]
    for col_i, h in enumerate(headers_diario, start=2):
        c = ws.cell(r, col_i, h)
        c.fill = BLUE_HEADER
        c.font = FONT_HEADER
        c.alignment = ALIGN_CENTER if col_i in (2, 3, 4) else (ALIGN_RIGHT if col_i in (6, 7) else ALIGN_LEFT)
    r += 1

    tot_debe = Decimal('0')
    tot_haber = Decimal('0')
    cuentas_map = {}

    for num_asiento, op in enumerate(operaciones_datos, 1):
        fecha_str = str(op.get('fecha') or '')
        glosa = str(op.get('glosa') or '')
        movs = op.get('movimientos', [])

        for idx, m in enumerate(movs):
            cod = str(m.get('codigo_cuenta') or '').strip()
            nom = str(m.get('nombre_cuenta') or f'Cuenta {cod}').strip()
            tipo = str(m.get('tipo') or '').strip().lower()
            try:
                monto = Decimal(str(m.get('monto', 0)).replace(',', '').strip())
            except Exception:
                monto = Decimal('0')

            if cod not in cuentas_map:
                cuentas_map[cod] = {'codigo': cod, 'nombre': nom, 'debe': [], 'haber': []}

            if tipo == 'debe':
                tot_debe += monto
                debe_val = float(monto)
                haber_val = None
                cuentas_map[cod]['debe'].append(monto)
            else:
                tot_haber += monto
                debe_val = None
                haber_val = float(monto)
                cuentas_map[cod]['haber'].append(monto)

            ws.cell(r, 2, num_asiento if idx == 0 else "").alignment = ALIGN_CENTER
            ws.cell(r, 3, fecha_str if idx == 0 else "").alignment = ALIGN_CENTER
            ws.cell(r, 4, cod).alignment = ALIGN_CENTER
            ws.cell(r, 5, nom).alignment = ALIGN_LEFT

            c_d = ws.cell(r, 6, debe_val)
            c_d.number_format = FORMAT_CURRENCY
            c_d.alignment = ALIGN_RIGHT

            c_h = ws.cell(r, 7, haber_val)
            c_h.number_format = FORMAT_CURRENCY
            c_h.alignment = ALIGN_RIGHT

            ws.cell(r, 8, glosa if idx == 0 else "").alignment = ALIGN_LEFT
            for col_i in range(2, 9):
                ws.cell(r, col_i).border = THIN_BORDER
                ws.cell(r, col_i).font = FONT_NORMAL
            r += 1

    # Fila de Totales del Diario
    ws.cell(r, 2, "TOTAL").alignment = ALIGN_CENTER
    ws.cell(r, 5, "SUMAS IGUALES DEL DIARIO").font = FONT_BOLD
    c_tot_d = ws.cell(r, 6, float(tot_debe))
    c_tot_d.number_format = FORMAT_CURRENCY
    c_tot_d.font = FONT_BOLD
    c_tot_d.alignment = ALIGN_RIGHT

    c_tot_h = ws.cell(r, 7, float(tot_haber))
    c_tot_h.number_format = FORMAT_CURRENCY
    c_tot_h.font = FONT_BOLD
    c_tot_h.alignment = ALIGN_RIGHT

    ws.cell(r, 8, "CUADRADO (Partida Doble)" if tot_debe == tot_haber else "DESCUADRADO").font = FONT_BOLD
    for col_i in range(2, 9):
        ws.cell(r, col_i).fill = TOTAL_FILL
        ws.cell(r, col_i).border = TOTAL_BORDER
    r += 3

    # ═════════════════════════════════════════════════════════════════════════
    # 2. REGISTRO EN LOS MAYORES (CUENTAS "T")
    # ═════════════════════════════════════════════════════════════════════════
    ws.cell(r, 2, "2. Registro en los mayores (Cuentas T):").font = Font(name="Calibri", size=12, bold=True, color="1F3C88")
    r += 1

    cuentas_ordenadas = sorted(cuentas_map.values(), key=lambda x: str(x['codigo']))
    idx_cuenta = 0

    while idx_cuenta < len(cuentas_ordenadas):
        c1 = cuentas_ordenadas[idx_cuenta]
        c2 = cuentas_ordenadas[idx_cuenta + 1] if idx_cuenta + 1 < len(cuentas_ordenadas) else None
        slots = [(c1, 2, 3)]
        if c2: slots.append((c2, 5, 6))

        fila_inicio_t = r
        max_filas = 0

        for c_data, col_d, col_h in slots:
            ws.merge_cells(start_row=fila_inicio_t, start_column=col_d, end_row=fila_inicio_t, end_column=col_h)
            tit_c = ws.cell(fila_inicio_t, col_d, f"{c_data['codigo']} {c_data['nombre'].upper()}")
            tit_c.font = Font(name="Calibri", size=11, bold=True, color="1F3C88")
            tit_c.alignment = ALIGN_CENTER

            cd = ws.cell(fila_inicio_t + 1, col_d, "D")
            cd.font = FONT_BOLD
            cd.alignment = ALIGN_CENTER
            cd.border = Border(bottom=Side(border_style="medium", color="000000"), right=Side(border_style="medium", color="000000"))

            ch = ws.cell(fila_inicio_t + 1, col_h, "H")
            ch.font = FONT_BOLD
            ch.alignment = ALIGN_CENTER
            ch.border = Border(bottom=Side(border_style="medium", color="000000"), left=Side(border_style="medium", color="000000"))

            movs_d = c_data['debe']
            movs_h = c_data['haber']
            n_filas = max(len(movs_d), len(movs_h), 1)
            max_filas = max(max_filas, n_filas)

            for i in range(n_filas):
                curr_r = fila_inicio_t + 2 + i
                c_izq = ws.cell(curr_r, col_d)
                c_der = ws.cell(curr_r, col_h)
                if i < len(movs_d):
                    c_izq.value = float(movs_d[i])
                    c_izq.number_format = FORMAT_CURRENCY
                    c_izq.alignment = ALIGN_RIGHT
                if i < len(movs_h):
                    c_der.value = float(movs_h[i])
                    c_der.number_format = FORMAT_CURRENCY
                    c_der.alignment = ALIGN_RIGHT
                c_izq.border = Border(right=Side(border_style="medium", color="000000"))
                c_der.border = Border(left=Side(border_style="medium", color="000000"))

        fila_saldo = fila_inicio_t + 2 + max_filas
        for c_data, col_d, col_h in slots:
            s_d = sum(c_data['debe'])
            s_h = sum(c_data['haber'])
            s_neto = s_d - s_h

            c_sd = ws.cell(fila_saldo, col_d)
            c_sh = ws.cell(fila_saldo, col_h)
            c_sd.border = Border(top=Side(border_style="medium", color="000000"), right=Side(border_style="medium", color="000000"))
            c_sh.border = Border(top=Side(border_style="medium", color="000000"), left=Side(border_style="medium", color="000000"))

            # Saldo neto destacado en amarillo según naturaleza
            cod_inicial = str(c_data['codigo']).strip()[:1]
            if cod_inicial in ('1', '2', '3', '6'):
                # Activo o Gasto (saldo normal Deudor)
                if s_neto >= 0:
                    c_sd.value = float(s_neto)
                    c_sd.number_format = FORMAT_CURRENCY
                    c_sd.font = FONT_BOLD
                    c_sd.fill = YELLOW_FILL
                    c_sd.alignment = ALIGN_RIGHT
                else:
                    c_sh.value = float(-s_neto)
                    c_sh.number_format = FORMAT_CURRENCY
                    c_sh.font = FONT_BOLD
                    c_sh.fill = YELLOW_FILL
                    c_sh.alignment = ALIGN_RIGHT
            else:
                # Pasivo, Patrimonio o Ingreso (saldo normal Acreedor)
                if s_neto <= 0:
                    c_sh.value = float(-s_neto)
                    c_sh.number_format = FORMAT_CURRENCY
                    c_sh.font = FONT_BOLD
                    c_sh.fill = YELLOW_FILL
                    c_sh.alignment = ALIGN_RIGHT
                else:
                    c_sd.value = float(s_neto)
                    c_sd.number_format = FORMAT_CURRENCY
                    c_sd.font = FONT_BOLD
                    c_sd.fill = YELLOW_FILL
                    c_sd.alignment = ALIGN_RIGHT

        r = fila_saldo + 2
        idx_cuenta += 2

    # ═════════════════════════════════════════════════════════════════════════
    # 3. ESTADOS FINANCIEROS (ESF Y ER) DINÁMICOS
    # ═════════════════════════════════════════════════════════════════════════
    ws.cell(r, 2, "3. Estados Financieros (ESF y ER):").font = Font(name="Calibri", size=12, bold=True, color="1F3C88")
    r += 1

    activos = []
    pasivos = []
    patrimonio = []
    ventas = Decimal('0')
    costo_ventas = Decimal('0')
    gastos_operativos = Decimal('0')
    otros_gastos = Decimal('0')

    for c in cuentas_ordenadas:
        cod = str(c['codigo']).strip()
        nom = c['nombre']
        s_d = sum(c['debe'])
        s_h = sum(c['haber'])
        pref = cod[:1]

        if pref in ('1', '2', '3'):
            saldo = s_d - s_h
            if saldo != 0: activos.append((cod, nom, saldo))
        elif pref == '4':
            saldo = s_h - s_d
            if saldo != 0: pasivos.append((cod, nom, saldo))
        elif pref == '5':
            saldo = s_h - s_d
            if saldo != 0: patrimonio.append((cod, nom, saldo))
        elif pref == '7':
            saldo = s_h - s_d
            ventas += saldo
        elif pref == '6':
            saldo = s_d - s_h
            if cod.startswith('69'):
                costo_ventas += saldo
            elif cod.startswith(('62', '63', '64', '65')):
                gastos_operativos += saldo
            else:
                otros_gastos += saldo

    utilidad_bruta = ventas - costo_ventas
    utilidad_operativa = utilidad_bruta - gastos_operativos
    utilidad_neta = utilidad_operativa - otros_gastos

    fila_ef = r
    # ── A. ESTADO DE SITUACIÓN FINANCIERA (Cols 2 a 5) ──
    ws.merge_cells(start_row=fila_ef, start_column=2, end_row=fila_ef, end_column=5)
    ws.cell(fila_ef, 2, "ESTADO DE SITUACION FINANCIERA").font = Font(name="Calibri", size=11, bold=True, color="1F3C88")
    ws.cell(fila_ef + 1, 2, "(Al cierre del período contable)").font = FONT_SUBTITLE

    r_esf = fila_ef + 3
    ws.cell(r_esf, 2, "ACTIVO").font = FONT_BOLD
    ws.cell(r_esf, 3, "IMPORTE").font = FONT_BOLD
    ws.cell(r_esf, 4, "PASIVO Y PATRIMONIO").font = FONT_BOLD
    ws.cell(r_esf, 5, "IMPORTE").font = FONT_BOLD
    for col_i in range(2, 6):
        ws.cell(r_esf, col_i).fill = TOTAL_FILL
        ws.cell(r_esf, col_i).border = THIN_BORDER
    r_esf += 1

    r_act = r_esf
    tot_act = Decimal('0')
    for cod, nom, s in activos:
        tot_act += s
        ws.cell(r_act, 2, f"{cod} {nom}").font = FONT_NORMAL
        c = ws.cell(r_act, 3, float(s))
        c.number_format = FORMAT_CURRENCY
        c.alignment = ALIGN_RIGHT
        r_act += 1

    ws.cell(r_act, 2, "TOTAL ACTIVO").font = FONT_BOLD
    c_ta = ws.cell(r_act, 3, float(tot_act))
    c_ta.font = FONT_BOLD
    c_ta.number_format = FORMAT_CURRENCY
    c_ta.alignment = ALIGN_RIGHT
    c_ta.fill = YELLOW_FILL

    r_pas = r_esf
    ws.cell(r_pas, 4, "PASIVO").font = Font(name="Calibri", size=10, bold=True, italic=True)
    r_pas += 1
    tot_pas = Decimal('0')
    for cod, nom, s in pasivos:
        tot_pas += s
        ws.cell(r_pas, 4, f"{cod} {nom}").font = FONT_NORMAL
        c = ws.cell(r_pas, 5, float(s))
        c.number_format = FORMAT_CURRENCY
        c.alignment = ALIGN_RIGHT
        r_pas += 1

    ws.cell(r_pas, 4, "TOTAL PASIVO").font = FONT_BOLD
    c = ws.cell(r_pas, 5, float(tot_pas))
    c.font = FONT_BOLD
    c.number_format = FORMAT_CURRENCY
    c.alignment = ALIGN_RIGHT
    r_pas += 2

    ws.cell(r_pas, 4, "PATRIMONIO").font = Font(name="Calibri", size=10, bold=True, italic=True)
    r_pas += 1
    tot_pat = Decimal('0')
    for cod, nom, s in patrimonio:
        tot_pat += s
        ws.cell(r_pas, 4, f"{cod} {nom}").font = FONT_NORMAL
        c = ws.cell(r_pas, 5, float(s))
        c.number_format = FORMAT_CURRENCY
        c.alignment = ALIGN_RIGHT
        r_pas += 1

    ws.cell(r_pas, 4, "RESULTADOS ACUMULADOS (Utilidad Neta)").font = FONT_BOLD
    c_res = ws.cell(r_pas, 5, float(utilidad_neta))
    c_res.font = FONT_BOLD
    c_res.number_format = FORMAT_CURRENCY
    c_res.alignment = ALIGN_RIGHT
    c_res.fill = GREEN_ACCENT
    r_pas += 1

    tot_pat_neto = tot_pat + utilidad_neta
    ws.cell(r_pas, 4, "TOTAL PATRIMONIO").font = FONT_BOLD
    c = ws.cell(r_pas, 5, float(tot_pat_neto))
    c.font = FONT_BOLD
    c.number_format = FORMAT_CURRENCY
    c.alignment = ALIGN_RIGHT
    r_pas += 1

    tot_pp = tot_pas + tot_pat_neto
    ws.cell(r_pas, 4, "TOTAL PASIVO + PATRIMONIO").font = FONT_BOLD
    c = ws.cell(r_pas, 5, float(tot_pp))
    c.font = FONT_BOLD
    c.number_format = FORMAT_CURRENCY
    c.alignment = ALIGN_RIGHT
    c.fill = YELLOW_FILL

    # ── B. ESTADO DE RESULTADOS (Cols 7 a 9) ──
    ws.merge_cells(start_row=fila_ef, start_column=7, end_row=fila_ef, end_column=9)
    ws.cell(fila_ef, 7, "ESTADO DE RESULTADOS").font = Font(name="Calibri", size=11, bold=True, color="1F3C88")
    ws.cell(fila_ef + 1, 7, "(Rendimiento del ejercicio)").font = FONT_SUBTITLE

    r_er = fila_ef + 3
    ws.cell(r_er, 7, "CONCEPTO").font = FONT_BOLD
    ws.cell(r_er, 8, "IMPORTE (S/)").font = FONT_BOLD
    for col_i in range(7, 9):
        ws.cell(r_er, col_i).fill = TOTAL_FILL
        ws.cell(r_er, col_i).border = THIN_BORDER
    r_er += 1

    lineas_er = [
        ("Ventas Netas", float(ventas), False, None),
        ("(-) Costo de Ventas", -float(costo_ventas), False, None),
        ("(=) UTILIDAD BRUTA", float(utilidad_bruta), True, TOTAL_FILL),
        ("(-) Gastos Operativos", -float(gastos_operativos), False, None),
        ("(=) UTILIDAD OPERATIVA", float(utilidad_operativa), True, TOTAL_FILL),
        ("(-) Otros Gastos", -float(otros_gastos), False, None),
        ("(=) UTILIDAD NETA DEL EJERCICIO", float(utilidad_neta), True, GREEN_ACCENT),
    ]

    for concepto, monto, es_bold, relleno in lineas_er:
        c1 = ws.cell(r_er, 7, concepto)
        c2 = ws.cell(r_er, 8, monto)
        c2.number_format = FORMAT_CURRENCY
        c2.alignment = ALIGN_RIGHT
        if es_bold:
            c1.font = FONT_BOLD
            c2.font = FONT_BOLD
        if relleno:
            c1.fill = relleno
            c2.fill = relleno
        c1.border = THIN_BORDER
        c2.border = THIN_BORDER
        r_er += 1

    ws.cell(r_er + 1, 7, "➡️ NOTA: La Utilidad Neta se traslada directamente a 'Resultados Acumulados' en el Patrimonio del ESF.").font = Font(name="Calibri", size=9, italic=True, color="1F3C88")

    ws.column_dimensions['B'].width = 10
    ws.column_dimensions['C'].width = 14
    ws.column_dimensions['D'].width = 14
    ws.column_dimensions['E'].width = 30
    ws.column_dimensions['F'].width = 16
    ws.column_dimensions['G'].width = 16
    ws.column_dimensions['H'].width = 30

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()


