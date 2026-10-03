"""Regresiones de carga, integración y saldos con el catálogo de clase."""
from datetime import date
from decimal import Decimal
from io import StringIO, BytesIO
import csv
import json
from unittest.mock import Mock, patch

from django.core.management import call_command
from django.db import IntegrityError, transaction
from django.test import TestCase
from django.template.loader import render_to_string
from django.urls import reverse

from .models import AsientoContable, CuentaContable, Movimiento
from .plan_cuentas import PLAN_CUENTAS
from .reporte_utils import get_reporte_context
from .services.chatbot import generar_propuesta, preparar_revision


class CargarPlanCuentasTests(TestCase):
    def cargar(self):
        salida, avisos = StringIO(), StringIO()
        call_command('cargar_plan_cuentas', stdout=salida, stderr=avisos)
        return salida.getvalue(), avisos.getvalue()

    def movimiento(self, codigo, tipo='debe', monto='100.00', asiento=None):
        if asiento is None:
            asiento = AsientoContable.objects.create(fecha=date(2026, 10, 2), descripcion='Original')
        return Movimiento.objects.create(
            asiento=asiento, cuenta=CuentaContable.objects.get(codigo=codigo),
            tipo=tipo, monto=Decimal(monto),
        )

    def snapshot(self):
        return (
            list(AsientoContable.objects.order_by('pk').values()),
            list(Movimiento.objects.order_by('pk').values()),
        )

    def test_carga_inicial_catalogo_exacto_sin_asientos(self):
        salida, avisos = self.cargar()
        codigos = {str(c) for c in (
            list(range(10, 19)) + list(range(20, 40)) + list(range(40, 50))
            + [50, 51, 52, 56, 57, 58] + list(range(60, 80))
        )} - {'15', '19'}
        self.assertEqual(CuentaContable.objects.count(), 64)
        self.assertSetEqual(set(CuentaContable.objects.values_list('codigo', flat=True)), codigos)
        self.assertEqual(len(PLAN_CUENTAS), 64)
        self.assertIn('Creadas: 64', salida)
        self.assertEqual(avisos, '')
        self.assertEqual(self.snapshot(), ([], []))

    def test_segunda_ejecucion_no_duplica_ni_actualiza(self):
        self.cargar()
        originales = list(CuentaContable.objects.values())
        salida, avisos = self.cargar()
        self.assertEqual(list(CuentaContable.objects.values()), originales)
        self.assertIn('Creadas: 0. Actualizadas: 0. Sin cambios: 64.', salida)
        self.assertEqual(avisos, '')

    def test_tipos_nombres_y_subcategorias_validos(self):
        self.cargar()
        esperados = {'activo': 28, 'pasivo': 10, 'patrimonio': 6, 'gasto': 10, 'ingreso': 10}
        for tipo, cantidad in esperados.items():
            self.assertEqual(CuentaContable.objects.filter(tipo=tipo).count(), cantidad)
        for codigo, nombre, tipo, subcategoria in PLAN_CUENTAS:
            with self.subTest(codigo=codigo):
                cuenta = CuentaContable.objects.get(codigo=codigo)
                cuenta.full_clean()
                self.assertEqual((cuenta.nombre, cuenta.tipo, cuenta.subcategoria), (nombre, tipo, subcategoria))
                self.assertIn(subcategoria, CuentaContable.SUBCATEGORIAS_POR_TIPO[tipo])
        self.assertEqual(CuentaContable.objects.get(codigo='29').tipo, 'activo')
        self.assertEqual(CuentaContable.objects.get(codigo='39').tipo, 'activo')
        self.assertFalse(CuentaContable.objects.filter(codigo__regex=r'^[89]').exists())

    def test_codigo_sigue_siendo_unico(self):
        self.cargar()
        with self.assertRaises(IntegrityError), transaction.atomic():
            CuentaContable.objects.create(codigo='11', nombre='Duplicada', tipo='activo')

    def test_conserva_cuentas_personalizadas_incluso_clases_excluidas(self):
        personalizada = CuentaContable.objects.create(
            codigo='94', nombre='Gastos Administrativos', tipo='gasto', subcategoria='gasto_operativo',
        )
        originales = list(CuentaContable.objects.filter(pk=personalizada.pk).values())
        self.cargar()
        self.cargar()
        self.assertEqual(list(CuentaContable.objects.filter(pk=personalizada.pk).values()), originales)
        self.assertEqual(CuentaContable.objects.count(), 65)

    def test_actualiza_nombre_y_clasificacion_sin_movimientos_conservando_id(self):
        cuenta = CuentaContable.objects.create(codigo='11', nombre='Antigua', tipo='pasivo')
        self.cargar()
        cuenta.refresh_from_db()
        self.assertEqual(cuenta.nombre, 'Inversiones Financieras')
        self.assertEqual(cuenta.tipo, 'activo')
        self.assertEqual(cuenta.subcategoria, 'activo_corriente')
        self.assertEqual(CuentaContable.objects.get(codigo='11').pk, cuenta.pk)

    def test_conserva_clasificacion_en_uso_y_avisa(self):
        cuenta = CuentaContable.objects.create(
            codigo='60', nombre='Compras antiguas', tipo='ingreso', subcategoria='otro_ingreso',
        )
        self.movimiento('60')
        originales = self.snapshot()
        _, avisos = self.cargar()
        cuenta.refresh_from_db()
        self.assertEqual(cuenta.nombre, 'Compras')
        self.assertEqual((cuenta.tipo, cuenta.subcategoria), ('ingreso', 'otro_ingreso'))
        self.assertIn('se conserva tipo=', avisos)
        self.assertIn('se conserva subcategoria=', avisos)
        self.assertEqual(self.snapshot(), originales)

    def test_conserva_subcategoria_de_gasto_en_uso_aunque_este_vacia(self):
        cuenta = CuentaContable.objects.create(codigo='69', nombre='Costo anterior', tipo='gasto')
        self.movimiento('69')
        _, avisos = self.cargar()
        cuenta.refresh_from_db()
        self.assertEqual(cuenta.subcategoria, '')
        self.assertIn("69: se conserva subcategoria=''", avisos)

    def test_conserva_asientos_movimientos_y_cuentas_originales(self):
        for codigo, nombre, tipo, subcat in [
            ('10', 'Efectivo y Equivalentes', 'activo', ''),
            ('20', 'Mercaderías', 'activo', ''),
            ('42', 'Cuentas por Pagar Comerciales', 'pasivo', ''),
            ('50', 'Capital Social', 'patrimonio', ''),
            ('60', 'Compras', 'gasto', 'costo_ventas'),
            ('70', 'Ventas', 'ingreso', ''),
        ]:
            CuentaContable.objects.create(codigo=codigo, nombre=nombre, tipo=tipo, subcategoria=subcat)
        asiento = AsientoContable.objects.create(fecha=date(2026, 10, 1), descripcion='Apertura')
        self.movimiento('10', asiento=asiento)
        self.movimiento('50', tipo='haber', asiento=asiento)
        originales = self.snapshot()
        ids = dict(CuentaContable.objects.values_list('codigo', 'pk'))
        self.cargar()
        self.cargar()
        self.assertEqual(self.snapshot(), originales)
        for codigo, pk in ids.items():
            self.assertEqual(CuentaContable.objects.get(codigo=codigo).pk, pk)
        self.assertEqual(CuentaContable.objects.count(), 64)
        self.assertTrue(AsientoContable.objects.get().esta_balanceado)

    def test_fallo_revierte_toda_la_carga(self):
        CuentaContable.objects.create(codigo='10', nombre='Original', tipo='activo')
        originales = list(CuentaContable.objects.values())
        crear = CuentaContable.objects.get_or_create

        def fallo(**kwargs):
            if kwargs['codigo'] == '18':
                raise IntegrityError('Fallo simulado')
            return crear(**kwargs)

        with patch('core.management.commands.cargar_plan_cuentas.CuentaContable.objects.get_or_create', side_effect=fallo):
            with self.assertRaises(IntegrityError):
                self.cargar()
        self.assertEqual(list(CuentaContable.objects.values()), originales)

    def test_plan_muestra_todas_las_cuentas_y_agrupaciones(self):
        self.cargar()
        pagina = self.client.get(reverse('gestionar_cuentas'))
        self.assertEqual(pagina.status_code, 200)
        self.assertEqual(len(pagina.context['cuentas']), 64)
        for cuenta in CuentaContable.objects.all():
            self.assertContains(pagina, cuenta.nombre)
        for etiqueta in ['Activo corriente', 'Existencias', 'Activo no corriente', 'Pasivo', 'Patrimonio', 'Gastos', 'Ingresos']:
            self.assertContains(pagina, etiqueta)

    def test_formulario_recibe_catalogo_completo_sin_movimientos(self):
        self.cargar()
        pagina = self.client.get(reverse('registrar_asiento'))
        self.assertEqual(pagina.status_code, 200)
        self.assertEqual(len(pagina.context['cuentas_json']), 64)
        self.assertSetEqual(
            {c['id'] for c in pagina.context['cuentas_json']},
            set(CuentaContable.objects.values_list('pk', flat=True)),
        )
        for codigo, nombre, _, _ in PLAN_CUENTAS:
            self.assertIn(f'{codigo} - {nombre}', [c['text'] for c in pagina.context['cuentas_json']])
        self.assertContains(pagina, 'Inversiones Financieras')

    def test_registra_asiento_con_cuentas_nuevas(self):
        self.cargar()
        respuesta = self.client.post(reverse('registrar_asiento'), {
            'fecha': '2026-10-02', 'descripcion': 'Obligación financiera', 'num_movimientos': '2',
            'cuenta_0': CuentaContable.objects.get(codigo='11').pk, 'tipo_0': 'debe', 'monto_0': '250.00',
            'cuenta_1': CuentaContable.objects.get(codigo='45').pk, 'tipo_1': 'haber', 'monto_1': '250.00',
        })
        self.assertEqual(respuesta.status_code, 302)
        asiento = AsientoContable.objects.get()
        self.assertTrue(asiento.esta_balanceado)
        self.assertSetEqual(set(asiento.movimientos.values_list('cuenta__codigo', flat=True)), {'11', '45'})

    def test_chatbot_recibe_catalogo_completo_y_resuelve_cuentas_nuevas(self):
        self.cargar()
        proveedor = Mock()
        proveedor.completar.return_value = json.dumps({
            'estado': 'propuesta', 'descripcion': 'Operación de prueba',
            'movimientos': [
                {'cuenta_codigo': '11', 'tipo': 'debe', 'monto': '100.00'},
                {'cuenta_codigo': '45', 'tipo': 'haber', 'monto': '100.00'},
            ],
        })
        originales = self.snapshot()
        propuesta = generar_propuesta('25/07/2020 | Operación de prueba', proveedor=proveedor)
        contexto = proveedor.completar.call_args.args[0][0]['content']
        catalogo = json.loads(contexto.split('CATÁLOGO COMPLETO:\n', 1)[1].split('\nFECHA DETECTADA:', 1)[0])
        self.assertEqual(len(catalogo), 64)
        self.assertEqual(catalogo, list(CuentaContable.objects.values('codigo', 'nombre', 'tipo', 'subcategoria')))
        revision = preparar_revision(propuesta)
        self.assertEqual(revision['movimientos'][0]['cuenta_id'], CuentaContable.objects.get(codigo='11').pk)
        self.assertEqual(self.snapshot(), originales)

    def test_formulario_plan_acepta_subcategorias_y_rechaza_incompatibles(self):
        for tipo, subcat in [('activo', 'existencias'), ('gasto', 'gasto_financiero'), ('ingreso', 'ingresos')]:
            with self.subTest(tipo=tipo):
                respuesta = self.client.post(reverse('gestionar_cuentas'), {
                    'codigo': f'custom-{tipo}', 'nombre': 'Personalizada', 'tipo': tipo, 'subcategoria': subcat,
                })
                self.assertEqual(respuesta.status_code, 302)
                self.assertEqual(CuentaContable.objects.get(codigo=f'custom-{tipo}').subcategoria, subcat)
        for subcat in ['desconocida', 'otro_ingreso']:
            respuesta = self.client.post(reverse('gestionar_cuentas'), {
                'codigo': 'incorrecta', 'nombre': 'Incorrecta', 'tipo': 'activo', 'subcategoria': subcat,
            })
            self.assertContains(respuesta, 'La subcategoría no es válida')
            self.assertFalse(CuentaContable.objects.filter(codigo='incorrecta').exists())


class ReportesPlanCuentasTests(TestCase):
    def cargar(self):
        call_command('cargar_plan_cuentas', stdout=StringIO(), stderr=StringIO())

    def asiento(self, debe, haber, monto):
        asiento = AsientoContable.objects.create(fecha=date(2026, 10, 2), descripcion='Prueba de reportes')
        for codigo, tipo in [(debe, 'debe'), (haber, 'haber')]:
            Movimiento.objects.create(
                asiento=asiento, cuenta=CuentaContable.objects.get(codigo=codigo), tipo=tipo, monto=Decimal(monto),
            )
        return asiento

    def contexto(self, vista):
        respuesta = self.client.get(reverse(vista))
        self.assertEqual(respuesta.status_code, 200)
        return respuesta.context

    def test_carga_no_altera_totales_de_los_cinco_reportes(self):
        for codigo, nombre, tipo, subcat in [
            ('10', 'Caja', 'activo', ''), ('20', 'Mercaderías', 'activo', ''),
            ('42', 'Proveedores', 'pasivo', ''), ('50', 'Capital Social', 'patrimonio', ''),
            ('60', 'Compras', 'gasto', 'costo_ventas'), ('70', 'Ventas', 'ingreso', ''),
        ]:
            CuentaContable.objects.create(codigo=codigo, nombre=nombre, tipo=tipo, subcategoria=subcat)
        self.asiento('10', '50', '1000')
        self.asiento('20', '42', '300')
        self.asiento('10', '70', '200')
        self.asiento('60', '20', '100')

        def snapshot_reportes():
            diario = self.contexto('libro_diario')
            mayor = self.contexto('libro_mayor')
            bc = self.contexto('balance_comprobacion')
            er = self.contexto('estado_resultados')
            bg = self.contexto('balance_general')
            exportado = get_reporte_context()
            return {
                'diario': [(a.pk, list(a.movimientos.values_list('pk', flat=True))) for a in diario['asientos']],
                'mayor': [(d['cuenta'].pk, d['total_debe'], d['total_haber'], d['saldo']) for d in mayor['datos_cuentas']],
                'bc': [bc[k] for k in ('gran_total_debe', 'gran_total_haber', 'gran_saldo_deudor', 'gran_saldo_acreedor', 'esta_cuadrado')],
                'er': [er[k] for k in ('total_ventas', 'total_costo_ventas', 'utilidad_bruta', 'utilidad_antes_impuesto', 'utilidad_neta')],
                'bg': [bg[k] for k in ('total_activos', 'total_pasivos', 'total_patrimonio', 'resultados_acumulados', 'esta_balanceado')],
                'exportado': {k: v for k, v in exportado.items() if isinstance(v, Decimal)},
            }

        originales = snapshot_reportes()
        self.cargar()
        self.cargar()
        self.assertEqual(snapshot_reportes(), originales)

    def test_correctoras_restan_del_activo_y_tienen_saldo_acreedor(self):
        self.cargar()
        self.asiento('20', '50', '1000')
        self.asiento('33', '50', '2000')
        for codigo in ['29', '36', '39']:
            self.asiento('68', codigo, '100')
        bg = self.contexto('balance_general')
        self.assertEqual(bg['total_activos'], Decimal('2700'))
        self.assertEqual(bg['resultados_acumulados'], Decimal('-300'))
        self.assertTrue(bg['esta_balanceado'])
        bc = self.contexto('balance_comprobacion')
        por_codigo = {d['cuenta'].codigo: d for d in bc['datos']}
        mayor = {d['cuenta'].codigo: d for d in self.contexto('libro_mayor')['datos_cuentas']}
        for codigo in ['29', '36', '39']:
            self.assertEqual(por_codigo[codigo]['saldo_acreedor'], Decimal('100'))
            self.assertEqual(mayor[codigo]['saldo'], Decimal('-100'))
        exportado = get_reporte_context()
        self.assertEqual(exportado['bg_total_activos'], bg['total_activos'])
        self.assertEqual(exportado['bg_total_pasivo_patrimonio'], bg['total_pasivo_patrimonio'])

    def test_resultados_y_exportaciones_incluyen_otros_ingresos_y_gastos(self):
        self.cargar()
        self.asiento('10', '50', '1000')
        self.asiento('10', '70', '500')
        self.asiento('69', '10', '100')
        self.asiento('62', '10', '50')
        self.asiento('67', '10', '20')
        self.asiento('10', '75', '40')
        self.asiento('66', '10', '10')
        # 74 conserva INGRESO y su débito reduce la utilidad sin nuevas reglas.
        self.asiento('74', '10', '5')
        er = self.contexto('estado_resultados')
        self.assertEqual(er['utilidad_antes_impuesto'], Decimal('355'))
        exportado = get_reporte_context()
        for clave in ['total_ventas', 'total_costo_ventas', 'total_gastos_operativos', 'total_gastos_financieros',
                      'total_otros_ingresos', 'total_otros_gastos', 'utilidad_antes_impuesto', 'utilidad_neta']:
            self.assertEqual(exportado[f'er_{clave}'], er[clave])
        bg = self.contexto('balance_general')
        self.assertTrue(bg['esta_balanceado'])
        self.assertEqual(exportado['bg_total_pasivo_patrimonio'], bg['total_pasivo_patrimonio'])
        mayor = {d['cuenta'].codigo: d['saldo_final'] for d in exportado['mayor_datos']}
        self.assertEqual(mayor['62'], Decimal('50'))
        self.assertEqual(mayor['74'], Decimal('-5'))

    def test_exportacion_conserva_signo_de_saldos_deudores_pasivo_patrimonio(self):
        self.cargar()
        self.asiento('42', '10', '100')
        self.asiento('50', '10', '50')
        bg = self.contexto('balance_general')
        exportado = get_reporte_context()
        self.assertEqual(bg['total_pasivos'], Decimal('-100'))
        self.assertEqual(bg['total_patrimonio'], Decimal('-50'))
        self.assertEqual(exportado['bg_total_pasivos'], bg['total_pasivos'])
        self.assertEqual(exportado['bg_total_patrimonio'], bg['total_patrimonio_con_resultados'])
        self.assertTrue(bg['esta_balanceado'])

    def test_csv_y_excel_desglosan_otros_resultados(self):
        from openpyxl import load_workbook

        self.cargar()
        self.asiento('10', '75', '100')
        self.asiento('66', '10', '25')
        csv_respuesta = self.client.get(reverse('exportar_reporte', args=['estado-resultados', 'csv']))
        self.assertEqual(csv_respuesta.status_code, 200)
        filas_csv = dict(list(csv.reader(StringIO(csv_respuesta.content.decode('utf-8-sig')), delimiter=';'))[1:])
        self.assertEqual(filas_csv['(+) Otros Ingresos'], '100.00')
        self.assertEqual(filas_csv['(-) Otros Gastos'], '-25.00')
        self.assertEqual(filas_csv['(=) RESULTADO ANTES DE IMPUESTO'], '75.00')
        excel = self.client.get(reverse('exportar_reporte', args=['estado-resultados', 'excel']))
        self.assertEqual(excel.status_code, 200)
        wb = load_workbook(BytesIO(excel.content))
        filas_excel = {c: m for c, m in wb['Estado de Resultados'].iter_rows(min_row=5, max_col=2, values_only=True)}
        self.assertEqual(filas_excel['(+) Otros Ingresos'], 100)
        self.assertEqual(filas_excel['(-) Otros Gastos'], -25)
        self.assertEqual(filas_excel['(=) RESULTADO ANTES DE IMPUESTOS'], 75)

    def test_plantilla_pdf_desglosa_otros_resultados(self):
        self.cargar()
        self.asiento('10', '75', '100')
        self.asiento('66', '10', '25')
        html = render_to_string('reporte_pdf.html', get_reporte_context())
        self.assertIn('(+) Otros Ingresos', html)
        self.assertIn('(-) Otros Gastos', html)
        self.assertIn('>100,00</td>', html)
        self.assertIn('>25,00</td>', html)
