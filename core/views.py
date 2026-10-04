"""
Vistas del Sistema Contable.

Cada vista implementa una funcionalidad específica:
- index: Menú principal con resumen
- gestionar_cuentas: CRUD del catálogo de cuentas contables
- registrar_asiento: Formulario dinámico para crear asientos
- libro_diario: Reporte de todos los asientos por fecha
- libro_mayor: Movimientos agrupados por cuenta con saldos
- balance_comprobacion: Verificación de cuadre (Debe == Haber)
- estado_resultados: Cálculo de utilidad/pérdida
- balance_general: Activo = Pasivo + Patrimonio
"""
import os
import requests
from PIL import Image, UnidentifiedImageError
import pytesseract
from dotenv import load_dotenv
import json
import time
import uuid
import gc
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib import messages
from django.db.models import Sum
from decimal import Decimal, InvalidOperation
from django.db import transaction
from django.conf import settings
from django.core.mail import EmailMessage
from django.template.loader import render_to_string
from django.http import HttpResponse, JsonResponse
from django.urls import reverse
from django.utils import timezone
from xhtml2pdf import pisa
from io import BytesIO
from .models import CuentaContable, AsientoContable, Movimiento
from .reporte_utils import get_reporte_context
from .services.asientos import validar_asiento_formulario
from .services.chatbot import ErrorChatbot, generar_propuesta, preparar_revision, texto_respuesta, informar_opciones
from .services.aclaracion_sesion import (CLAVE_ACLARACION, obtener_pendiente, guardar_pendiente,
                                        nuevo_hecho_completo, resolver_respuesta, consulta_opciones)


# ─── PÁGINA PRINCIPAL ──────────────────────────────────────────────────────────

def index(request):
    """
    Muestra el menú principal con un resumen del estado del sistema:
    total de cuentas, asientos, y totales de Debe/Haber.
    """
    total_asientos = AsientoContable.objects.count()
    total_cuentas = CuentaContable.objects.count()
    total_movimientos = Movimiento.objects.count()

    # Calcular totales generales
    total_debe = Movimiento.objects.filter(tipo='debe').aggregate(
        total=Sum('monto'))['total'] or Decimal('0')
    total_haber = Movimiento.objects.filter(tipo='haber').aggregate(
        total=Sum('monto'))['total'] or Decimal('0')

    # Últimos 5 asientos registrados (ordenados por el momento en que se crearon en el sistema)
    ultimos_asientos = AsientoContable.objects.prefetch_related(
        'movimientos__cuenta'
    ).order_by('-created_at')[:5]

    context = {
        'total_asientos': total_asientos,
        'total_cuentas': total_cuentas,
        'total_movimientos': total_movimientos,
        'total_debe': total_debe,
        'total_haber': total_haber,
        'ultimos_asientos': ultimos_asientos,
    }
    return render(request, 'index.html', context)


# ─── GESTIÓN DE CUENTAS CONTABLES ──────────────────────────────────────────────

def gestionar_cuentas(request):
    """
    Permite crear nuevas cuentas contables especificando código, nombre, tipo
    y subcategoría (para clasificación en el Estado de Resultados).
    También muestra la lista completa del catálogo de cuentas.
    """
    if request.method == 'POST':
        codigo = request.POST.get('codigo', '').strip()
        nombre = request.POST.get('nombre', '').strip()
        tipo = request.POST.get('tipo', '').strip()
        subcategoria = request.POST.get('subcategoria', '').strip()

        errors = []
        if not codigo:
            errors.append('El código es obligatorio.')
        if not nombre:
            errors.append('El nombre es obligatorio.')
        if tipo not in dict(CuentaContable.TIPO_CHOICES):
            errors.append('El tipo de cuenta no es válido.')
        if subcategoria and subcategoria not in CuentaContable.SUBCATEGORIAS_POR_TIPO.get(tipo, []):
            errors.append('La subcategoría no es válida para este tipo de cuenta.')
        if CuentaContable.objects.filter(codigo=codigo).exists():
            errors.append(f'Ya existe una cuenta con el código "{codigo}".')

        if errors:
            for error in errors:
                messages.error(request, error)
        else:
            CuentaContable.objects.create(
                codigo=codigo, nombre=nombre, tipo=tipo, subcategoria=subcategoria
            )
            messages.success(request, f'Cuenta "{codigo} - {nombre}" creada exitosamente.')
            return redirect('gestionar_cuentas')

    cuentas = CuentaContable.objects.all()
    context = {
        'cuentas': cuentas,
        'tipos': CuentaContable.TIPO_CHOICES,
        'subcategorias': CuentaContable.SUBCATEGORIA_CHOICES,
        'subcategorias_por_tipo': CuentaContable.SUBCATEGORIAS_POR_TIPO,
        'formulario': request.POST if request.method == 'POST' else {},
    }
    return render(request, 'gestionar_cuentas.html', context)


def eliminar_cuenta(request, cuenta_id):
    """
    Elimina una cuenta contable solo si no tiene movimientos asociados.
    Si tiene movimientos, muestra un mensaje de error.
    """
    cuenta = get_object_or_404(CuentaContable, id=cuenta_id)
    if cuenta.movimientos.exists():
        messages.error(
            request,
            f'No se puede eliminar "{cuenta}" porque tiene movimientos asociados.'
        )
    else:
        messages.success(request, f'Cuenta "{cuenta}" eliminada exitosamente.')
        cuenta.delete()
    return redirect('gestionar_cuentas')


# ─── REGISTRO DE ASIENTOS CONTABLES ────────────────────────────────────────────

PROPUESTAS_SESION = 'chatbot_propuestas'
PROPUESTA_TTL = 30 * 60


def _propuestas_vigentes(request):
    ahora = time.time()
    return {
        identificador: propuesta
        for identificador, propuesta in request.session.get(PROPUESTAS_SESION, {}).items()
        if 0 <= ahora - propuesta['creada_en'] < PROPUESTA_TTL
    }


def _formulario_enviado(datos):
    try:
        cantidad = max(0, min(int(datos.get('num_movimientos', 0)), 1000))
    except (ValueError, TypeError):
        cantidad = 0
    return {
        'fecha': datos.get('fecha', ''), 'descripcion': datos.get('descripcion', ''),
        'movimientos': [
            {'cuenta_id': datos.get(f'cuenta_{i}', ''), 'tipo': datos.get(f'tipo_{i}', ''),
             'monto': datos.get(f'monto_{i}', '')}
            for i in range(cantidad)
        ],
    }


def registrar_asiento(request):
    """
    Formulario dinámico para registrar asientos contables.
    Procesa múltiples líneas de movimiento (mínimo 2).
    Valida que la suma del Debe sea igual a la del Haber antes de guardar.
    """
    cuentas = CuentaContable.objects.all()

    if not cuentas.exists():
        messages.warning(
            request,
            'Debe crear al menos una cuenta contable antes de registrar asientos.'
        )
        return redirect('gestionar_cuentas')

    inicial = {'fecha': '', 'descripcion': '', 'movimientos': []}
    propuesta_id = request.GET.get('propuesta', '')
    if request.method == 'GET' and propuesta_id:
        propuestas = _propuestas_vigentes(request)
        propuesta = propuestas.get(propuesta_id)
        if propuesta is None:
            messages.warning(request, 'La propuesta no existe en esta sesión o ha caducado.')
            return redirect('registrar_asiento')
        try:
            revision = preparar_revision(propuesta['resultado'])
        except ErrorChatbot as exc:
            propuestas.pop(propuesta_id, None)
            request.session[PROPUESTAS_SESION] = propuestas
            messages.error(request, str(exc) + ' ' + ' '.join(exc.errores))
            return redirect('registrar_asiento')
        inicial.update(fecha=revision['fecha'], descripcion=revision['descripcion'], movimientos=revision['movimientos'])

    if request.method == 'POST':
        inicial = _formulario_enviado(request.POST)
        fecha = request.POST.get('fecha', '').strip()
        descripcion = request.POST.get('descripcion', '').strip()
        movimientos_data, errors = validar_asiento_formulario(request.POST)

        if errors:
            for error in errors:
                messages.error(request, error)
        else:
            with transaction.atomic():
                # Crear el asiento y sus movimientos en la base de datos
                asiento = AsientoContable.objects.create(
                    fecha=fecha,
                    descripcion=descripcion,
                )
                for m in movimientos_data:
                    Movimiento.objects.create(
                        asiento=asiento,
                        cuenta=m['cuenta'],
                        tipo=m['tipo'],
                        monto=m['monto'],
                    )

            if propuesta_id:
                propuestas = _propuestas_vigentes(request)
                propuestas.pop(propuesta_id, None)
                request.session[PROPUESTAS_SESION] = propuestas

            # --- MENSAJE CON HORA REAL ---
            hora_real = timezone.localtime(timezone.now()).strftime("%d/%m/%Y a las %H:%M")
            messages.success(
                request,
                f'Asiento registrado exitosamente el {hora_real}.'
            )
            return redirect('libro_diario')

    # Para el formulario mostramos los últimos creados
    asientos_registrados = AsientoContable.objects.prefetch_related('movimientos__cuenta').order_by('-created_at')

    context = {
        'cuentas': cuentas,
        'asientos_registrados': asientos_registrados,
        'formulario_inicial': inicial,
        'propuesta_id': propuesta_id,
        'cuentas_json': [{'id': c.pk, 'codigo': c.codigo, 'nombre': c.nombre,
                         'text': f'{c.codigo} - {c.nombre}'} for c in cuentas],
    }
    return render(request, 'registrar_asiento.html', context)


# ─── LIBRO DIARIO ──────────────────────────────────────────────────────────────

def libro_diario(request):
    """
    Muestra todos los asientos contables ordenados cronológicamente.
    Cada asiento se despliega con sus movimientos al Debe y al Haber.
    """
    # --- ORDEN CRONOLÓGICO ESTRICTO ---
    asientos = AsientoContable.objects.prefetch_related('movimientos__cuenta').order_by('fecha', 'id')
    cuentas = CuentaContable.objects.all()
    
    context = {
        'asientos': asientos,
        'cuentas': cuentas,
    }
    return render(request, 'libro_diario.html', context)


# ─── LIBRO MAYOR ───────────────────────────────────────────────────────────────

def libro_mayor(request):
    """
    Agrupa los movimientos por cuenta contable.
    Para cada cuenta muestra:
    - Todos sus movimientos con fecha y asiento
    - Total Debe, Total Haber
    - Saldo calculado según la naturaleza de la cuenta
    """
    cuentas = CuentaContable.objects.all()

    datos_cuentas = []
    for cuenta in cuentas:
        # ORDEN CRONOLÓGICO EN EL MAYOR
        movimientos = cuenta.movimientos.select_related('asiento').order_by('asiento__fecha', 'asiento__id')
        if not movimientos.exists():
            continue

        total_debe = movimientos.filter(tipo='debe').aggregate(
            total=Sum('monto'))['total'] or Decimal('0')
        total_haber = movimientos.filter(tipo='haber').aggregate(
            total=Sum('monto'))['total'] or Decimal('0')

        # Saldo según naturaleza de la cuenta
        if cuenta.naturaleza_deudora:
            saldo = total_debe - total_haber
        else:
            saldo = total_haber - total_debe

        datos_cuentas.append({
            'cuenta': cuenta,
            'movimientos': movimientos,
            'total_debe': total_debe,
            'total_haber': total_haber,
            'saldo': saldo,
        })

    context = {
        'datos_cuentas': datos_cuentas,
    }
    return render(request, 'libro_mayor.html', context)


# ─── BALANCE DE COMPROBACIÓN ───────────────────────────────────────────────────

def balance_comprobacion(request):
    """
    Lista todas las cuentas con movimientos y verifica que los totales
    de Debe y Haber coincidan (el balance esté cuadrado).
    
    Muestra columnas de Sumas (Debe/Haber) y Saldos (Deudor/Acreedor).
    """
    cuentas = CuentaContable.objects.all()

    datos = []
    gran_total_debe = Decimal('0')
    gran_total_haber = Decimal('0')
    gran_saldo_deudor = Decimal('0')
    gran_saldo_acreedor = Decimal('0')

    for cuenta in cuentas:
        total_debe = cuenta.movimientos.filter(tipo='debe').aggregate(
            total=Sum('monto'))['total'] or Decimal('0')
        total_haber = cuenta.movimientos.filter(tipo='haber').aggregate(
            total=Sum('monto'))['total'] or Decimal('0')

        if total_debe == 0 and total_haber == 0:
            continue

        saldo = total_debe - total_haber
        saldo_deudor = saldo if saldo > 0 else Decimal('0')
        saldo_acreedor = abs(saldo) if saldo < 0 else Decimal('0')

        datos.append({
            'cuenta': cuenta,
            'total_debe': total_debe,
            'total_haber': total_haber,
            'saldo_deudor': saldo_deudor,
            'saldo_acreedor': saldo_acreedor,
        })

        gran_total_debe += total_debe
        gran_total_haber += total_haber
        gran_saldo_deudor += saldo_deudor
        gran_saldo_acreedor += saldo_acreedor

    context = {
        'datos': datos,
        'gran_total_debe': gran_total_debe,
        'gran_total_haber': gran_total_haber,
        'gran_saldo_deudor': gran_saldo_deudor,
        'gran_saldo_acreedor': gran_saldo_acreedor,
        'esta_cuadrado': gran_total_debe == gran_total_haber,
    }
    return render(request, 'balance_comprobacion.html', context)


# ─── ESTADO DE RESULTADOS (DETALLADO) ──────────────────────────────────────────

def estado_resultados(request):
    """
    Estado de Resultados con desglose completo.
    """
    def saldo_cuenta(cuenta):
        t_debe = cuenta.movimientos.filter(tipo='debe').aggregate(
            total=Sum('monto'))['total'] or Decimal('0')
        t_haber = cuenta.movimientos.filter(tipo='haber').aggregate(
            total=Sum('monto'))['total'] or Decimal('0')
        if cuenta.tipo == 'gasto':
            return t_debe - t_haber
        else:  # ingreso
            return t_haber - t_debe

    def obtener_items(queryset):
        items = []
        total = Decimal('0')
        for c in queryset:
            s = saldo_cuenta(c)
            if s != 0:
                items.append({'cuenta': c, 'saldo': s})
                total += s
        return items, total

    ventas, total_ventas = obtener_items(
        CuentaContable.objects.filter(tipo='ingreso').exclude(subcategoria='otro_ingreso')
    )
    costo_ventas, total_costo_ventas = obtener_items(
        CuentaContable.objects.filter(tipo='gasto', subcategoria='costo_ventas')
    )
    utilidad_bruta = total_ventas - total_costo_ventas

    gastos_operativos, total_gastos_operativos = obtener_items(
        CuentaContable.objects.filter(tipo='gasto').exclude(
            subcategoria__in=['costo_ventas', 'gasto_financiero', 'otro_gasto']
        )
    )
    utilidad_operativa = utilidad_bruta - total_gastos_operativos

    gastos_financieros, total_gastos_financieros = obtener_items(
        CuentaContable.objects.filter(tipo='gasto', subcategoria='gasto_financiero')
    )
    otros_ingresos, total_otros_ingresos = obtener_items(
        CuentaContable.objects.filter(tipo='ingreso', subcategoria='otro_ingreso')
    )
    otros_gastos, total_otros_gastos = obtener_items(
        CuentaContable.objects.filter(tipo='gasto', subcategoria='otro_gasto')
    )

    utilidad_antes_impuesto = (
        utilidad_operativa
        - total_gastos_financieros
        + total_otros_ingresos
        - total_otros_gastos
    )

    tasa_impuesto = Decimal('30')
    if utilidad_antes_impuesto > 0:
        impuesto = (utilidad_antes_impuesto * tasa_impuesto / Decimal('100')).quantize(Decimal('0.01'))
    else:
        impuesto = Decimal('0')

    utilidad_neta = utilidad_antes_impuesto - impuesto

    context = {
        'ventas': ventas, 'total_ventas': total_ventas,
        'costo_ventas': costo_ventas, 'total_costo_ventas': total_costo_ventas,
        'utilidad_bruta': utilidad_bruta,
        'gastos_operativos': gastos_operativos, 'total_gastos_operativos': total_gastos_operativos,
        'utilidad_operativa': utilidad_operativa,
        'gastos_financieros': gastos_financieros, 'total_gastos_financieros': total_gastos_financieros,
        'otros_ingresos': otros_ingresos, 'total_otros_ingresos': total_otros_ingresos,
        'otros_gastos': otros_gastos, 'total_otros_gastos': total_otros_gastos,
        'utilidad_antes_impuesto': utilidad_antes_impuesto,
        'tasa_impuesto': tasa_impuesto,
        'impuesto': impuesto,
        'utilidad_neta': utilidad_neta,
    }
    return render(request, 'estado_resultados.html', context)


# ─── BALANCE GENERAL (ESTADO DE SITUACIÓN FINANCIERA) ──────────────────────────

def balance_general(request):
    """
    Estado de Situación Financiera con verificación de la ecuación contable.
    """
    def calcular_saldos_por_tipo(tipo_cuenta, deudora=True):
        cuentas = CuentaContable.objects.filter(tipo=tipo_cuenta)
        items = []
        total = Decimal('0')
        for cuenta in cuentas:
            t_debe = cuenta.movimientos.filter(tipo='debe').aggregate(
                total=Sum('monto'))['total'] or Decimal('0')
            t_haber = cuenta.movimientos.filter(tipo='haber').aggregate(
                total=Sum('monto'))['total'] or Decimal('0')
            saldo = (t_debe - t_haber) if deudora else (t_haber - t_debe)
            if saldo != 0:
                items.append({'cuenta': cuenta, 'saldo': saldo})
                total += saldo
        return items, total

    activos, total_activos = calcular_saldos_por_tipo('activo', deudora=True)
    pasivos, total_pasivos = calcular_saldos_por_tipo('pasivo', deudora=False)
    patrimonio, total_patrimonio = calcular_saldos_por_tipo('patrimonio', deudora=False)

    _, total_ingresos = calcular_saldos_por_tipo('ingreso', deudora=False)
    _, total_gastos = calcular_saldos_por_tipo('gasto', deudora=True)
    resultados_acumulados = total_ingresos - total_gastos  

    total_patrimonio_con_resultados = total_patrimonio + resultados_acumulados
    total_pasivo_patrimonio = total_pasivos + total_patrimonio_con_resultados
    esta_balanceado = total_activos == total_pasivo_patrimonio

    def totales_debe_haber(tipo_cuenta):
        cuentas = CuentaContable.objects.filter(tipo=tipo_cuenta)
        total_d = Decimal('0')
        total_h = Decimal('0')
        for c in cuentas:
            td = c.movimientos.filter(tipo='debe').aggregate(
                total=Sum('monto'))['total'] or Decimal('0')
            th = c.movimientos.filter(tipo='haber').aggregate(
                total=Sum('monto'))['total'] or Decimal('0')
            total_d += td
            total_h += th
        return total_d, total_h

    activo_debe, activo_haber = totales_debe_haber('activo')
    pasivo_debe, pasivo_haber = totales_debe_haber('pasivo')
    patrimonio_debe, patrimonio_haber = totales_debe_haber('patrimonio')

    ecuacion_resumen = [
        {
            'concepto': 'ACTIVO',
            'debe': activo_debe, 'haber': activo_haber,
            'saldo_deudor': activo_debe - activo_haber if activo_debe >= activo_haber else Decimal('0'),
            'saldo_acreedor': activo_haber - activo_debe if activo_haber > activo_debe else Decimal('0'),
        },
        {
            'concepto': 'PASIVO',
            'debe': pasivo_debe, 'haber': pasivo_haber,
            'saldo_deudor': pasivo_debe - pasivo_haber if pasivo_debe >= pasivo_haber else Decimal('0'),
            'saldo_acreedor': pasivo_haber - pasivo_debe if pasivo_haber > pasivo_debe else Decimal('0'),
        },
        {
            'concepto': 'PATRIMONIO',
            'debe': patrimonio_debe, 'haber': patrimonio_haber,
            'saldo_deudor': patrimonio_debe - patrimonio_haber if patrimonio_debe >= patrimonio_haber else Decimal('0'),
            'saldo_acreedor': patrimonio_haber - patrimonio_debe if patrimonio_haber > patrimonio_debe else Decimal('0'),
        },
    ]
    ec_total_debe = sum(e['debe'] for e in ecuacion_resumen)
    ec_total_haber = sum(e['haber'] for e in ecuacion_resumen)
    ec_total_deudor = sum(e['saldo_deudor'] for e in ecuacion_resumen)
    ec_total_acreedor = sum(e['saldo_acreedor'] for e in ecuacion_resumen)

    context = {
        'activos': activos, 'pasivos': pasivos, 'patrimonio': patrimonio,
        'total_activos': total_activos, 'total_pasivos': total_pasivos,
        'total_patrimonio': total_patrimonio, 'resultados_acumulados': resultados_acumulados,
        'es_utilidad': resultados_acumulados >= 0,
        'total_patrimonio_con_resultados': total_patrimonio_con_resultados,
        'total_pasivo_patrimonio': total_pasivo_patrimonio, 'esta_balanceado': esta_balanceado,
        'ecuacion_resumen': ecuacion_resumen, 'ec_total_debe': ec_total_debe,
        'ec_total_haber': ec_total_haber, 'ec_total_deudor': ec_total_deudor,
        'ec_total_acreedor': ec_total_acreedor,
    }
    return render(request, 'balance_general.html', context)


# ─── REPORTE COMPLETO EN PDF ───────────────────────────────────────────────────

import gc # Recolector de basura
from django.utils import timezone

def reporte_completo(request):
    if request.method == 'POST':
        empresa = request.POST.get('empresa', 'Mi Empresa').strip() or 'Mi Empresa'

        # Liberamos RAM antes de empezar
        gc.collect()

        context = get_reporte_context()
        context['empresa'] = empresa
        
        # --- LOGO EN BASE64 ---
        import base64
        try:
            logo_path = os.path.join(settings.BASE_DIR, 'UNILOGO.png')
            with open(logo_path, "rb") as image_file:
                encoded_string = base64.b64encode(image_file.read()).decode('utf-8')
                context['logo_base64'] = f"data:image/png;base64,{encoded_string}"
        except Exception:
            context['logo_base64'] = ""

        html = render_to_string('reporte_pdf.html', context)
        del context
        
        pdf_file = BytesIO()
        pisa_status = pisa.CreatePDF(BytesIO(html.encode('UTF-8')), dest=pdf_file)
        del html
        gc.collect()

        if pisa_status.err:
            messages.error(request, 'Error al generar el archivo PDF.')
            return redirect('reporte_completo')

        # Descarga directa del PDF en el navegador
        pdf_data = pdf_file.getvalue()
        pdf_file.close()

        nombre_archivo = f"Reporte_Contable_{empresa.replace(' ', '_')}.pdf"
        response = HttpResponse(pdf_data, content_type='application/pdf')
        response['Content-Disposition'] = f'attachment; filename="{nombre_archivo}"'
        return response

    return render(request, 'menu_reporte.html')

def eliminar_asiento(request, asiento_id):
    """
    Elimina un asiento contable específico y todos sus movimientos asociados.
    """
    asiento = get_object_or_404(AsientoContable, id=asiento_id)
    
    # Guardamos el ID para el mensaje de éxito antes de borrarlo
    asiento_num = asiento.id 
    
    # Al eliminar el asiento, los movimientos se borran en cascada 
    # por el models.CASCADE que tienes en tu modelo
    asiento.delete()
    
    messages.success(request, f'El Asiento #{asiento_num} fue eliminado correctamente.')
    return redirect('libro_diario')


# --- AÑADIR AL FINAL DE core/views.py ---

def editar_asiento(request, asiento_id):
    asiento = get_object_or_404(AsientoContable, id=asiento_id)
    if request.method == 'POST':
        movimientos_data, errors = validar_asiento_formulario(request.POST)
        if errors:
            for error in errors:
                messages.error(request, error)
            return redirect('libro_diario')

        with transaction.atomic():
            asiento.fecha = request.POST.get('fecha', '').strip()
            asiento.descripcion = request.POST.get('descripcion', '').strip()
            asiento.save()
            asiento.movimientos.all().delete()
            for movimiento in movimientos_data:
                Movimiento.objects.create(asiento=asiento, **movimiento)

        # --- MENSAJE CON HORA REAL ---
        ahora = timezone.localtime(timezone.now()).strftime("%d/%m/%Y a las %H:%M")
        messages.success(request, f"Asiento actualizado exitosamente el {ahora}.")
    return redirect('libro_diario')
# ─── CHATBOT DE ASISTENCIA FINANCIERA (GROQ AI) ──────────────────────────

def chatbot_api(request):
    """Propone asientos validados sin guardar, conservando el texto del chat."""
    if request.method != 'POST':
        return JsonResponse({'error': 'Método no permitido'}, status=405)

    if request.content_type == 'application/json':
        try:
            datos = json.loads(request.body)
        except (ValueError, UnicodeDecodeError):
            return JsonResponse({'error': 'JSON de solicitud no válido'}, status=400)
        if not isinstance(datos, dict):
            return JsonResponse({'error': 'JSON de solicitud no válido'}, status=400)
        mensaje = datos.get('message', '')
    else:
        mensaje = request.POST.get('message', '')
    if not isinstance(mensaje, str) or not mensaje.strip():
        return JsonResponse({'error': 'Mensaje vacío o no válido'}, status=400)

    try:
        mensaje = mensaje.strip()
        pendiente = obtener_pendiente(request.session)
        if mensaje.lower().strip('.!') in {'cancelar', 'cancela', 'cancelar aclaración', 'cancelar aclaracion'}:
            request.session.pop(CLAVE_ACLARACION, None)
            return JsonResponse({'estado': 'cancelado', 'response': 'Aclaración cancelada.', 'status': 'success'})
        cuentas = list(CuentaContable.objects.all())
        if pendiente and consulta_opciones(mensaje):
            resultado = informar_opciones(mensaje, pendiente, cuentas)
            campo = pendiente.get('campo_pendiente', pendiente['datos_faltantes'][0])
            if any(p in campo for p in ('cuenta', 'tipo_gasto', 'contrapartida')):
                pendiente = {**pendiente, 'opciones_validas': [
                    o if 'id' in o else
                    {'id': o['cuenta_codigo'], 'etiqueta': o['cuenta_nombre'], 'valor': o}
                    for o in resultado['opciones']]}
                pendiente['opciones'] = pendiente['opciones_validas']
            request.session[CLAVE_ACLARACION] = {**pendiente, 'actualizada_en': time.time()}
            return JsonResponse({**resultado, 'status': 'success'})
        if pendiente and nuevo_hecho_completo(mensaje, cuentas):
            request.session.pop(CLAVE_ACLARACION, None)
            pendiente = None
        original = pendiente['mensaje_original'] if pendiente else mensaje
        if pendiente:
            hecho, detectados, seguimiento, _ = resolver_respuesta(mensaje, pendiente, cuentas)
            resultado = seguimiento.get('aclaracion_resolver') or generar_propuesta(
                hecho, datos_previos=detectados, seguimiento=seguimiento)
        else:
            resultado = generar_propuesta(mensaje)
        if resultado['estado'] == 'requiere_aclaracion':
            guardar_pendiente(request.session, original, resultado, pendiente, mensaje)
        if resultado['estado'] == 'propuesta':
            revision = preparar_revision(resultado)
            request.session.pop(CLAVE_ACLARACION, None)
            propuestas = _propuestas_vigentes(request)
            # Acotar datos temporales de la sesión; no almacenar conversaciones.
            propuestas = dict(list(propuestas.items())[-4:])
            identificador = uuid.uuid4().hex
            propuestas[identificador] = {'resultado': resultado, 'creada_en': time.time()}
            request.session[PROPUESTAS_SESION] = propuestas
            resultado = {
                **resultado, 'vista_previa': revision, 'propuesta_id': identificador,
                'revision_url': reverse('registrar_asiento') + '?propuesta=' + identificador,
            }

    except ErrorChatbot as exc:
        error = {'status': 'error', 'error': str(exc)}
        if exc.errores:
            error['errores'] = exc.errores
            error['error'] += ' ' + ' '.join(exc.errores)
        return JsonResponse(error, status=exc.status)
    return JsonResponse({**resultado, 'response': texto_respuesta(resultado), 'status': 'success'})


# ─── EXPORTACIÓN A EXCEL Y CSV ────────────────────────────────────────────────

def exportar_reporte(request, reporte, formato):
    """
    Controlador para exportar reportes a Excel (.xlsx) o CSV.
    reporte: 'completo', 'diario', 'balance-comprobacion', 'estado-resultados', 'balance-general'
    formato: 'excel', 'csv'
    """
    from . import export_utils
    from openpyxl import Workbook
    import io

    formato = formato.lower()
    reporte = reporte.lower()

    if formato == 'excel':
        wb = Workbook()
        if reporte in ('caso-hoja-unica', 'caso', 'caso-completo'):
            content = export_utils.generar_excel_caso_hoja_unica()
            filename = "Caso_Ciclo_Contable_UNI.xlsx"
        elif reporte == 'completo':
            content = export_utils.generar_excel_completo()
            filename = "Reporte_Contable_Completo.xlsx"
        elif reporte == 'diario':
            export_utils.build_excel_diario(wb, title_sheet="Libro Diario")
            buf = io.BytesIO()
            wb.save(buf)
            content = buf.getvalue()
            filename = "Libro_Diario.xlsx"
        elif reporte == 'balance-comprobacion':
            export_utils.build_excel_balance_comprobacion(wb, title_sheet="Balance Comprobación")
            if "Sheet" in wb.sheetnames and len(wb.sheetnames) > 1:
                del wb["Sheet"]
            buf = io.BytesIO()
            wb.save(buf)
            content = buf.getvalue()
            filename = "Balance_de_Comprobacion.xlsx"
        elif reporte == 'estado-resultados':
            export_utils.build_excel_estado_resultados(wb, title_sheet="Estado de Resultados")
            if "Sheet" in wb.sheetnames:
                del wb["Sheet"]
            buf = io.BytesIO()
            wb.save(buf)
            content = buf.getvalue()
            filename = "Estado_de_Resultados.xlsx"
        elif reporte == 'balance-general':
            export_utils.build_excel_balance_general(wb, title_sheet="Balance General")
            if "Sheet" in wb.sheetnames:
                del wb["Sheet"]
            buf = io.BytesIO()
            wb.save(buf)
            content = buf.getvalue()
            filename = "Balance_General.xlsx"
        else:
            return HttpResponse("Tipo de reporte no válido.", status=400)

        response = HttpResponse(
            content,
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response

    elif formato == 'csv':
        if reporte == 'diario':
            content = export_utils.generar_csv_diario()
            filename = "Libro_Diario.csv"
        elif reporte == 'balance-comprobacion':
            content = export_utils.generar_csv_balance_comprobacion()
            filename = "Balance_de_Comprobacion.csv"
        elif reporte == 'estado-resultados':
            content = export_utils.generar_csv_estado_resultados()
            filename = "Estado_de_Resultados.csv"
        elif reporte == 'balance-general':
            content = export_utils.generar_csv_balance_general()
            filename = "Balance_General.csv"
        else:
            return HttpResponse("Reporte CSV no disponible para esta opción.", status=400)

        response = HttpResponse(content, content_type='text/csv; charset=utf-8-sig')
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response

    return HttpResponse("Formato no soportado.", status=400)


# ─── OCR Y ESTRUCTURACIÓN DE ASIENTOS CON IA ─────────────────────────────────

def _validar_cuentas_ocr(operaciones, catalogo):
    """Verifica la estructura y resuelve únicamente códigos exactos del catálogo."""
    por_codigo = {}
    for cuenta in catalogo:
        por_codigo.setdefault(str(cuenta['codigo']).strip(), []).append(cuenta)

    for operacion in operaciones:
        if not isinstance(operacion, dict) or not isinstance(operacion.get('movimientos'), list):
            raise ValueError('Operación OCR inválida.')
        for movimiento in operacion['movimientos']:
            if not isinstance(movimiento, dict):
                raise ValueError('Movimiento OCR inválido.')
            tipo = str(movimiento.get('tipo', '')).strip().lower()
            if tipo not in ('debe', 'haber'):
                raise ValueError('Movimiento OCR inválido.')
            movimiento['tipo'] = tipo
            try:
                raw_monto = str(movimiento.get('monto', '')).strip().replace(',', '')
                monto = Decimal(raw_monto)
            except (InvalidOperation, TypeError) as exc:
                raise ValueError('Monto OCR inválido.') from exc
            if not monto.is_finite() or monto <= 0:
                raise ValueError('Monto OCR inválido.')

            codigo = movimiento.get('codigo_cuenta')
            codigo = str(codigo).strip() if codigo is not None else ''
            coincidencias = por_codigo.get(codigo, []) if codigo else []
            pendiente = movimiento.get('pendiente_revision', False)
            if len(coincidencias) == 1 and pendiente is False:
                cuenta = coincidencias[0]
                movimiento.update(codigo_cuenta=cuenta['codigo'], nombre_cuenta=cuenta['nombre'],
                                  pendiente_revision=False)
            else:
                # No propagar un código inventado, ambiguo o marcado como incierto.
                movimiento.update(codigo_cuenta=None, nombre_cuenta='', pendiente_revision=True,
                                  motivo_revision='Cuenta no resuelta con seguridad en el catálogo. '
                                                  'Seleccione una cuenta manualmente.')


def procesar_imagen_asiento(request):
    """
    Procesa una imagen de comprobante u operación contable mediante OCR (pytesseract)
    y utiliza la IA de Groq para proponer asientos con el catálogo registrado.
    """
    if request.method != 'POST':
        return JsonResponse({'error': 'Método no permitido. Utilice POST.'}, status=405)

    # Aceptar múltiples imágenes subidas o texto directo
    imagenes = request.FILES.getlist('imagenes') or request.FILES.getlist('imagen')
    texto_directo = request.POST.get('texto', '').strip()

    texto_ocr = ""

    if imagenes:
        # Configuración de ruta tesseract en Windows si no está en PATH
        if os.path.exists(r'C:\Program Files\Tesseract-OCR\tesseract.exe'):
            pytesseract.pytesseract.tesseract_cmd = r'C:\Program Files\Tesseract-OCR\tesseract.exe'
        elif os.path.exists(r'C:\Program Files (x86)\Tesseract-OCR\tesseract.exe'):
            pytesseract.pytesseract.tesseract_cmd = r'C:\Program Files (x86)\Tesseract-OCR\tesseract.exe'

        try:
            # 1. Extracción de texto con OCR para cada imagen recibida
            textos_extraidos = []
            for idx, img_file in enumerate(imagenes):
                with Image.open(img_file) as img:
                    try:
                        txt = pytesseract.image_to_string(img, lang='spa+eng', timeout=30)
                    except pytesseract.TesseractError:
                        txt = pytesseract.image_to_string(img, timeout=30)
                txt = txt.strip()
                if txt:
                    textos_extraidos.append(f"--- IMAGEN/DOCUMENTO {idx + 1} ---\n{txt}")

            texto_ocr = "\n\n".join(textos_extraidos).strip()
        except UnidentifiedImageError:
            return JsonResponse({'error': 'El archivo enviado no es una imagen válida.'}, status=400)
        except pytesseract.TesseractNotFoundError:
            return JsonResponse({'error': 'Tesseract OCR no está instalado o no se encuentra en PATH.'}, status=503)
        except Exception as ocr_err:
            return JsonResponse({'error': f'Error en el motor OCR: {str(ocr_err)}'}, status=500)
    elif texto_directo:
        texto_ocr = texto_directo
    else:
        return JsonResponse({'error': 'No se proporcionó ningún archivo de imagen ni texto.'}, status=400)

    if not texto_ocr:
        return JsonResponse({'error': 'No se detectó texto legible en las imágenes subidas.'}, status=400)

    try:
        # 2. Configurar llamada a Groq API (recargando .env dinámicamente)
        load_dotenv(os.path.join(settings.BASE_DIR, '.env'), override=True)
        api_key = os.environ.get('GROQ_API_KEY')
        if not api_key:
            return JsonResponse({
                'error': 'Falta configurar GROQ_API_KEY en el archivo .env para el procesamiento con IA.',
                'texto_ocr_detectado': texto_ocr
            }, status=500)

        catalogo = list(CuentaContable.objects.values('codigo', 'nombre', 'tipo', 'subcategoria'))
        system_prompt = (
            "Eres un contador profesional y docente de contabilidad financiera (PCGE - Plan Contable General Empresarial). "
            "Tu misión es analizar exhaustivamente TODO el texto u operaciones del caso contable y estructurar la lista COMPLETA de todas las operaciones del ciclo contable en formato JSON.\n\n"
            "Estructura JSON estrictamente requerida:\n"
            "{\n"
            '  "titulo_caso": "Nombre descriptivo del caso (ej: Empresa ABC SAC - Periodo X)",\n'
            '  "operaciones": [\n'
            "    {\n"
            '      "fecha": "YYYY-MM-DD",\n'
            '      "glosa": "Descripción clara del hecho",\n'
            '      "movimientos": [\n'
            "        {\n"
            '          "codigo_cuenta": "codigo_del_catalogo",\n'
            '          "nombre_cuenta": "nombre_del_catalogo",\n'
            '          "tipo": "debe_o_haber",\n'
            '          "monto": 1000.00,\n'
            '          "pendiente_revision": false\n'
            "        }\n"
            "      ]\n"
            "    }\n"
            "  ]\n"
            "}\n\n"
            "CRITERIOS CONTABLES PARA RESOLVER CUALQUIER CASO X:\n"
            "1. EXHAUSTIVIDAD TOTAL: Si el texto o imagen contiene varias diapositivas, hojas o partes apiladas verticalmente (por ejemplo, con títulos como 'Continuación', o Diapositiva 19 y 20), debes extraer e incluir TODAS las operaciones de TODAS las partes hasta el final (hasta 'Se pide'). NUNCA te detengas en la mitad.\n"
            "2. PARTIDA DOBLE OBLIGATORIA: En cada operación, la suma exacta del 'debe' debe ser idéntica a la suma del 'haber'.\n"
            "3. CATÁLOGO ESTRICTO: Usa ÚNICAMENTE códigos existentes en el catálogo real proporcionado como strings. No inventes códigos ni devuelvas códigos que no existan en este catálogo. Si alguna cuenta no se puede determinar con total seguridad, márcala con pendiente_revision: true.\n"
            "4. FECHAS REALES: Identifica las fechas reales del enunciado (formato YYYY-MM-DD). Si no indica día específico, asígnales fechas secuenciales coherentes dentro del mes del ejercicio.\n"
            "5. TRATAMIENTO DE OPERACIONES TÍPICAS:\n"
            "   - Aporte inicial o constitución: Activos aportados (10 Efectivo, 20 Mercaderías, 33 IME) al 'debe', contra 50 Capital al 'haber'.\n"
            "   - Compra de mercaderías y activos fijos: 20 Mercaderías o 33 Inmuebles, Maquinaria y Equipo al 'debe', contra 10 Efectivo o 42/46 Cuentas por Pagar al 'haber'.\n"
            "   - Venta de mercaderías: 10 Efectivo y/o 12 Cuentas por Cobrar al 'debe', contra 70 Ventas al 'haber'.\n"
            "   - Devolución de mercadería con reembolso: 10 Efectivo al 'debe' / 20 Mercaderías al 'haber'.\n"
            "   - Cobro de factura / cuentas por cobrar: 10 Efectivo al 'debe' / 12 Cuentas por Cobrar al 'haber'.\n"
            "   - Donación de mercadería recibida: 20 Mercaderías al 'debe' / 75 Otros Ingresos de Gestión al 'haber'.\n"
            "   - Pérdida por siniestro o incendio: 65 Otros Gastos de Gestión al 'debe' / 20 Mercaderías al 'haber'.\n"
            "   - Gastos devengados impagos (sueldos pendientes): 62 Gastos de Personal al 'debe' / 41 Remuneraciones por Pagar al 'haber'.\n"
            "   - Gastos de servicios (alquiler, luz, agua): 63 Gastos de Servicios Prestados por Terceros al 'debe' / 10 Efectivo al 'haber'.\n"
            "6. CÁLCULO DINÁMICO DE COSTO DE VENTAS:\n"
            "   - Si el caso menciona inventario final (conteo físico de mercaderías al cierre): calcula: "
            "Costo de Ventas = Inventario Inicial + Compras de mercaderías (y donaciones) - Devoluciones - Pérdida por incendio - Inventario Final. "
            "Genera la operación al cierre: glosa 'Ajuste de costo de ventas por inventario final', con 69 Costo de Ventas al 'debe' y 20 Mercaderías al 'haber'.\n"
            "7. AJUSTE DE DEPRECIACIÓN:\n"
            "   - Si se adquirieron activos fijos con vida útil estimada y se pide preparar Estados Financieros al cierre: "
            "calcula la depreciación del período transcurrido: (Costo - Valor de rescate) / Vida útil. "
            "Genera el asiento al cierre: 68 Valuación y Deterioro al 'debe' / 39 Depreciación Acumulada al 'haber'.\n"
            "8. Devuelve ÚNICAMENTE el objeto JSON sin explicaciones adicionales ni código markdown.\n\n"
            "Catálogo real de cuentas (JSON):\n" + json.dumps(catalogo, ensure_ascii=False)
        )

        url = "https://api.groq.com/openai/v1/chat/completions"
        headers = {
            "Authorization": f"Bearer {api_key}",
            "Content-Type": "application/json"
        }

        # Modelos activos en Groq
        modelos = ["openai/gpt-oss-120b", "qwen/qwen3.8-27b", "openai/gpt-oss-20b", "llama-3.3-70b-versatile"]
        ai_content = None
        last_err = None

        for modelo in modelos:
            try:
                payload = {
                    "model": modelo,
                    "messages": [
                        {"role": "system", "content": system_prompt},
                        {"role": "user", "content": f"Texto de operaciones contables:\n\n{texto_ocr}"}
                    ],
                    "response_format": {"type": "json_object"},
                    "max_tokens": 3500,
                    "temperature": 0.1
                }
                response = requests.post(url, json=payload, headers=headers, timeout=30)
                if response.status_code == 200:
                    ai_content = response.json()['choices'][0]['message']['content']
                    if ai_content:
                        break
                else:
                    last_err = f"{response.status_code}: {response.text}"
            except Exception as e:
                last_err = str(e)
                continue

        if not ai_content:
            return JsonResponse({'error': 'No se pudo estructurar el texto con IA. Puede revisar el texto OCR.', 'texto_ocr_detectado': texto_ocr}, status=500)

        # 3. Parsear y limpiar respuesta JSON
        cleaned_json = ai_content.strip()
        if cleaned_json.startswith("```json"):
            cleaned_json = cleaned_json[7:]
        if cleaned_json.startswith("```"):
            cleaned_json = cleaned_json[3:]
        if cleaned_json.endswith("```"):
            cleaned_json = cleaned_json[:-3]
        cleaned_json = cleaned_json.strip()

        data = json.loads(cleaned_json)
        if not isinstance(data, dict) or not isinstance(data.get('operaciones'), list):
            return JsonResponse({'error': 'La IA no devolvió una lista de operaciones válida.',
                                 'texto_ocr_detectado': texto_ocr}, status=502)
        try:
            _validar_cuentas_ocr(data['operaciones'], catalogo)
        except ValueError:
            return JsonResponse({'error': 'La IA devolvió movimientos inválidos.',
                                 'texto_ocr_detectado': texto_ocr}, status=502)
        data['texto_ocr_detectado'] = texto_ocr
        return JsonResponse(data)

    except json.JSONDecodeError:
        return JsonResponse({
            'error': 'La IA no devolvió una estructura JSON válida.',
            'texto_ocr_detectado': texto_ocr
        }, status=500)
    except Exception as e:
        return JsonResponse({'error': 'No se pudo estructurar el texto con IA.', 'texto_ocr_detectado': texto_ocr}, status=500)


def transcribir_audio(request):
    """
    Recibe un archivo de audio grabado desde el micrófono y lo transcribe
    usando la API de Groq Whisper (whisper-large-v3-turbo).
    """
    if request.method != 'POST':
        return JsonResponse({'error': 'Método no permitido. Use POST.'}, status=405)

    audio_file = request.FILES.get('audio')
    if not audio_file:
        return JsonResponse({'error': 'No se recibió ningún archivo de audio.'}, status=400)

    load_dotenv(os.path.join(settings.BASE_DIR, '.env'), override=True)
    api_key = os.environ.get('GROQ_API_KEY')
    if not api_key:
        return JsonResponse({'error': 'Falta configurar GROQ_API_KEY en el archivo .env.'}, status=500)

    url = 'https://api.groq.com/openai/v1/audio/transcriptions'
    headers = {'Authorization': f'Bearer {api_key}'}
    nombre = getattr(audio_file, 'name', 'audio.webm') or 'audio.webm'
    files = {
        'file': (nombre, audio_file.read(), audio_file.content_type or 'audio/webm')
    }
    data = {
        'model': 'whisper-large-v3-turbo',
        'language': 'es',
        'response_format': 'json'
    }

    try:
        response = requests.post(url, headers=headers, files=files, data=data, timeout=30)
        if response.status_code == 200:
            resultado = response.json()
            return JsonResponse({'texto': resultado.get('text', '').strip()})
        else:
            # Fallback a whisper-large-v3
            data['model'] = 'whisper-large-v3'
            response2 = requests.post(url, headers=headers, files=files, data=data, timeout=30)
            if response2.status_code == 200:
                resultado = response2.json()
                return JsonResponse({'texto': resultado.get('text', '').strip()})
    except Exception as e:
        return JsonResponse({'error': f'Error al conectar con el servicio de voz: {str(e)}'}, status=500)


def guardar_operaciones_lote(request):
    """
    Guarda en la base de datos en una sola transacción atómica todas las
    operaciones detectadas por la IA / OCR.
    """
    if request.method != 'POST':
        return JsonResponse({'error': 'Método no permitido. Use POST.'}, status=405)

    try:
        data = json.loads(request.body)
        operaciones = data.get('operaciones', [])
        if not operaciones:
            return JsonResponse({'error': 'No se enviaron operaciones para guardar.'}, status=400)

        catalogo = {str(c.codigo).strip(): c for c in CuentaContable.objects.all()}
        guardados = 0

        with transaction.atomic():
            for op in operaciones:
                fecha = op.get('fecha') or timezone.now().strftime('%Y-%m-%d')
                glosa = (op.get('glosa') or 'Asiento registrado por IA').strip()
                movs = op.get('movimientos', [])
                if not movs:
                    continue

                asiento = AsientoContable.objects.create(fecha=fecha, descripcion=glosa)
                for m in movs:
                    codigo = str(m.get('codigo_cuenta') or '').strip()
                    cuenta = catalogo.get(codigo)
                    if not cuenta:
                        raise ValueError(f"La cuenta con código '{codigo}' no existe en el catálogo.")

                    tipo = str(m.get('tipo', '')).strip().lower()
                    raw_monto = str(m.get('monto', 0)).replace(',', '').strip()
                    monto = Decimal(raw_monto)
                    Movimiento.objects.create(asiento=asiento, cuenta=cuenta, tipo=tipo, monto=monto)
                guardados += 1

        return JsonResponse({
            'status': 'success',
            'guardados': guardados,
            'mensaje': f'¡Se guardaron exitosamente los {guardados} asientos del caso en la base de datos!'
        })
    except Exception as e:
        return JsonResponse({'error': f'Error al guardar operaciones: {str(e)}'}, status=400)


def exportar_caso_ia_excel(request):
    """
    Genera y descarga el archivo Excel de hoja única (Diario, Cuentas T, ESF, ER)
    directamente a partir de las operaciones enviadas por la IA (POST) o de la base de datos (GET).
    Permite resolver y descargar de inmediato cualquier Caso X.
    """
    from . import export_utils

    operaciones = None
    titulo_caso = "CASO CICLO CONTABLE"

    if request.method == 'POST':
        try:
            if request.content_type == 'application/json':
                datos = json.loads(request.body)
            else:
                raw = request.POST.get('datos_caso', '')
                datos = json.loads(raw) if raw else {}

            operaciones = datos.get('operaciones')
            titulo_caso = datos.get('titulo_caso') or datos.get('titulo') or titulo_caso
        except Exception:
            operaciones = None
    elif request.method == 'GET':
        titulo_caso = request.GET.get('titulo', titulo_caso)

    content = export_utils.generar_excel_caso_hoja_unica(
        operaciones_datos=operaciones,
        titulo_caso=titulo_caso
    )

    nombre_limpio = "".join(c for c in titulo_caso if c.isalnum() or c in (' ', '_', '-')).strip()
    nombre_archivo = f"{nombre_limpio.replace(' ', '_')}.xlsx" if nombre_limpio else "Caso_Ciclo_Contable.xlsx"

    response = HttpResponse(
        content,
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{nombre_archivo}"'
    return response


def limpiar_asientos_caso(request):
    """
    Elimina todos los asientos registrados en la base de datos para iniciar un caso nuevo desde cero.
    """
    if request.method != 'POST':
        return JsonResponse({'error': 'Método no permitido. Use POST.'}, status=405)

    try:
        with transaction.atomic():
            total = AsientoContable.objects.count()
            AsientoContable.objects.all().delete()

        return JsonResponse({
            'status': 'success',
            'mensaje': f'Se han eliminado los {total} asientos anteriores. La base de datos está limpia para resolver un nuevo caso.'
        })
    except Exception as e:
        return JsonResponse({'error': f'Error al limpiar asientos: {str(e)}'}, status=500)



